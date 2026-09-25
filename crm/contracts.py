from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .models import Contract


TEMPLATE_DIRECTORY = Path(__file__).resolve().parent / "contract_templates"
TEMPLATE_FILES = {
    Contract.Kind.CLIENT_FORWARDING: "client_forwarding.docx",
    Contract.Kind.CARRIER_TRANSPORT: "carrier_transport.docx",
    Contract.Kind.SUBCONTRACTOR_FORWARDING: "subcontractor_forwarding.docx",
}

MONTHS = (
    "",
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)


def russian_date(value):
    return f"«{value.day:02d}» {MONTHS[value.month]} {value.year} г."


def short_date(value):
    return value.strftime("%d.%m.%Y")


def value_or_dash(value):
    return str(value).strip() if value else "—"


def _capitalize_like(value, source):
    return value.capitalize() if source[:1].isupper() else value


def _genitive_word(word, *, surname=False):
    """Small deterministic declension for full Russian names in contract headers."""
    lowered = word.lower()
    if len(word) < 3 or "-" in word:
        return word
    if lowered.endswith(("ов", "ев", "ёв", "ин", "ын", "ой")):
        if lowered.endswith("ой"):
            return _capitalize_like(word[:-2] + "ого", word)
        return _capitalize_like(word + "а", word)
    if surname and lowered.endswith(("ова", "ева", "ёва", "ина", "ына")):
        return _capitalize_like(word[:-1] + "ой", word)
    if lowered.endswith("ский"):
        return _capitalize_like(word[:-2] + "ого", word)
    if lowered.endswith("цкий"):
        return _capitalize_like(word[:-2] + "ого", word)
    if lowered.endswith(("ич", "вич")):
        return _capitalize_like(word + "а", word)
    if lowered.endswith(("на", "вна")):
        return _capitalize_like(word[:-1] + "ы", word)
    if lowered.endswith("ий"):
        return _capitalize_like(word[:-2] + "ия", word)
    if lowered.endswith("ей"):
        return _capitalize_like(word[:-2] + "ея", word)
    if lowered.endswith("а"):
        return _capitalize_like(word[:-1] + "ы", word)
    if lowered.endswith("я"):
        return _capitalize_like(word[:-1] + "и", word)
    if lowered.endswith(("н", "р", "л", "м", "т", "д", "б", "г", "в", "п", "к", "х", "ч", "ш", "щ", "ж")):
        return _capitalize_like(word + "а", word)
    return word


def genitive_full_name(value):
    """Return a presentation-ready genitive full name; keeps incomplete input intact."""
    pieces = [piece for piece in str(value or "").strip().split() if piece]
    if len(pieces) < 2:
        return " ".join(pieces)
    return " ".join(
        _genitive_word(piece, surname=index == 0)
        for index, piece in enumerate(pieces)
    )


def genitive_position(value):
    position = " ".join(str(value or "").split())
    known = {
        "генеральный директор": "Генерального директора",
        "директор": "Директора",
        "коммерческий директор": "Коммерческого директора",
        "исполнительный директор": "Исполнительного директора",
        "индивидуальный предприниматель": "Индивидуального предпринимателя",
    }
    return known.get(position.lower(), position)


def _party_organization(party):
    return getattr(party, "organization", None)


def _is_entrepreneur(party):
    organization = _party_organization(party)
    if getattr(organization, "kind", "") == "entrepreneur":
        return True
    name = str(getattr(party, "name", "")).strip().lower()
    return name.startswith("индивидуальный предприниматель") or name.startswith("ип ")


def _representative_text(party, name, position):
    organization = _party_organization(party)
    display_name = value_or_dash(
        name
        or getattr(party, "director_name", "")
        or getattr(organization, "director_name", "")
    )
    display_position = (
        position
        or getattr(party, "director_position", "")
        or getattr(organization, "director_position", "")
        or "Генеральный директор"
    )
    gender_feminine = any(
        part.lower().endswith(("вна", "ична", "на"))
        for part in str(display_name).split()
    )
    acting = "действующей" if gender_feminine else "действующего"
    if _is_entrepreneur(party):
        return f"{genitive_full_name(display_name)}, {acting}"
    return f"{genitive_position(display_position)} {genitive_full_name(display_name)}, {acting}"


def _authority_basis(contract, prefix, party):
    authority_type = getattr(contract, f"{prefix}_authority_type", "")
    basis = str(getattr(contract, f"{prefix}_authority_basis", "") or "").strip()
    if _is_entrepreneur(party) and authority_type == Contract.AuthorityType.CHARTER:
        return "листа записи ЕГРИП"
    if authority_type == Contract.AuthorityType.CHARTER or basis.lower() in {"устав", "устава"}:
        return "Устава"
    normalized = {
        "доверенность": "доверенности",
        "доверенности": "доверенности",
        "приказ": "приказа",
        "решение": "решения",
    }
    return normalized.get(basis.lower(), basis or "документа, подтверждающего полномочия")


def party_address(party):
    return getattr(party, "legal_address", "") or getattr(party, "address", "")


def party_requisites(heading, party, representative):
    inn_kpp = value_or_dash(getattr(party, "tax_id", ""))
    if getattr(party, "kpp", ""):
        inn_kpp += f" / {party.kpp}"
    return "\n".join(
        (
            heading,
            value_or_dash(party.name),
            f"ИНН/КПП: {inn_kpp}",
            f"ОГРН: {value_or_dash(getattr(party, 'ogrn', ''))}",
            f"Адрес: {value_or_dash(party_address(party))}",
            f"р/с: {value_or_dash(getattr(party, 'settlement_account', ''))}",
            f"Банк: {value_or_dash(getattr(party, 'bank_name', ''))}",
            f"БИК: {value_or_dash(getattr(party, 'bik', ''))}",
            f"к/с: {value_or_dash(getattr(party, 'correspondent_account', ''))}",
            f"E-mail: {value_or_dash(getattr(party, 'email', ''))}",
            "",
            f"________________ / {value_or_dash(representative)} /",
        )
    )


def _remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def _set_paragraph(paragraph, text, bold=False):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.bold = bold


def _contract_roles(contract):
    if contract.kind == Contract.Kind.CLIENT_FORWARDING:
        return {
            "left": "ЭКСПЕДИТОР",
            "right": "КЛИЕНТ",
            "left_role": "Экспедитор",
            "right_role": "Клиент",
            "subtitle": (
                f"{contract.expeditor} — Экспедитор / "
                f"{contract.counterparty_name} — Клиент"
            ),
        }
    if contract.kind == Contract.Kind.CARRIER_TRANSPORT:
        return {
            "left": "ЗАКАЗЧИК",
            "right": "ПЕРЕВОЗЧИК",
            "left_role": "Заказчик",
            "right_role": "Перевозчик",
            "subtitle": (
                f"{contract.expeditor} — Заказчик / "
                f"{contract.counterparty_name} — Перевозчик"
            ),
        }
    return {
        "left": "КЛИЕНТ",
        "right": "ЭКСПЕДИТОР",
        "left_role": "Клиент",
        "right_role": "Экспедитор",
        "subtitle": (
            f"{contract.expeditor} — Клиент / "
            f"{contract.counterparty_name} — привлечённый Экспедитор"
        ),
    }


def _intro_text(contract, roles):
    left_representative = _representative_text(
        contract.expeditor,
        contract.expeditor_representative,
        contract.expeditor_representative_position,
    )
    right_representative = _representative_text(
        contract.counterparty,
        contract.counterparty_representative,
        contract.counterparty_representative_position,
    )
    return (
        f"{contract.expeditor.name}, именуемое в дальнейшем «{roles['left_role']}», "
        f"в лице {left_representative} на основании "
        f"{_authority_basis(contract, 'expeditor', contract.expeditor)}, с одной стороны, и "
        f"{contract.counterparty_name}, именуемое в дальнейшем "
        f"«{roles['right_role']}», в лице {right_representative} на "
        f"основании {_authority_basis(contract, 'counterparty', contract.counterparty)}, с другой "
        "стороны, совместно именуемые «Стороны», заключили настоящий Договор."
    )


def build_contract_docx(contract):
    template_path = TEMPLATE_DIRECTORY / TEMPLATE_FILES[contract.kind]
    document = Document(template_path)
    roles = _contract_roles(contract)

    title = document.paragraphs[0]
    title_text = re.sub(r"№\s*[^\n]+$", f"№ {contract.number}", title.text)
    _set_paragraph(title, title_text, bold=True)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if len(document.paragraphs) > 2:
        _set_paragraph(document.paragraphs[2], f"г. {contract.city}\t\t{russian_date(contract.contract_date)}")
    if len(document.paragraphs) > 4:
        _set_paragraph(document.paragraphs[4], _intro_text(contract, roles))

    if document.tables:
        requisites_table = document.tables[0]
        requisites_table.cell(0, 0).text = party_requisites(
            roles["left"], contract.expeditor, contract.expeditor_representative
        )
        requisites_table.cell(0, 1).text = party_requisites(
            roles["right"], contract.counterparty, contract.counterparty_representative
        )

    if contract.valid_until:
        for paragraph in document.paragraphs:
            if "Договор действует" in paragraph.text or "Договор вступает" in paragraph.text:
                text = re.sub(
                    r"до\s+(?:«?_+»?\s+_+\s+20_+\s*г\.|_+)",
                    f"до {short_date(contract.valid_until)}",
                    paragraph.text,
                )
                text = re.sub(
                    r"Договор действует[^.]*\.",
                    f"Договор действует до {short_date(contract.valid_until)}.",
                    text,
                    count=1,
                )
                _set_paragraph(paragraph, text)

    # Служебное примечание из клиентского шаблона не является частью договора.
    for paragraph in list(document.paragraphs):
        if paragraph.text.strip().startswith("Примечание к шаблону") or paragraph.text.strip().startswith(
            "Шаблон подготовлен как коммерческий"
        ):
            _remove_paragraph(paragraph)

    properties = document.core_properties
    properties.title = f"Договор № {contract.number}"
    properties.subject = contract.get_kind_display()
    properties.comments = "Сформирован в Экспедитор CRM из утверждённого шаблона."

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream
