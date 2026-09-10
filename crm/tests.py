from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
import json
import tempfile
from unittest.mock import patch
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document

from .forms import (
    ShipmentForm,
    TransportationChainForm,
    TransportationDocumentForm,
    VehicleCombinationForm,
)
from .epd import epd_validation_errors, prepare_documents
from .models import (
    BankStatement,
    BankStatementLine,
    Carrier,
    ChatMessage,
    CompanyProfile,
    Contract,
    Customer,
    DocumentBatch,
    DocumentBatchLine,
    Driver,
    DriverEmployment,
    DriverLicense,
    DriverPassport,
    DirectConversation,
    ForwardingOrder,
    Organization,
    OrganizationBankAccount,
    OrganizationContact,
    OrganizationGroup,
    OrganizationRole,
    Payment,
    PlannerTask,
    ReconciliationAct,
    ReconciliationActLine,
    SettlementMovement,
    Shipment,
    ShipmentDocument,
    TransportOrder,
    TransportOrderStop,
    Transportation,
    TransportationInstruction,
    TransportationLink,
    TransportationParty,
    TransportationStatusEvent,
    TransportationElectronicDocument,
    TripCharge,
    TransportationStop,
    UserProfile,
    VATRate,
    VehicleAssignment,
    VehicleCombination,
    Vehicle,
)


class CrmTestCase(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="manager", password="test-password"
        )
        self.user.crm_profile.role = UserProfile.Role.ADMIN
        self.user.crm_profile.save(update_fields=["role", "updated_at"])
        self.company_profile = CompanyProfile.objects.create(
            name='ООО "Тестовый экспедитор"',
            short_name='ООО "Тестовый экспедитор"',
            tax_id="7701000000",
            kpp="770101001",
            legal_address="Москва, ул. Экспедиторская, 1",
            bank_name='АО "Тест Банк"',
            bik="044525000",
            settlement_account="40702810000000000001",
            correspondent_account="30101810000000000000",
            director_name="Петрова Анна Сергеевна",
            chief_accountant_name="Петрова Анна Сергеевна",
            default_vat_rate=CompanyProfile.VATRate.TWENTY_TWO,
        )
        self.customer = Customer.objects.create(
            name="Тестовый клиент",
            tax_id="7711000000",
            kpp="771101001",
            ogrn="1027700000001",
            address="Москва, ул. Клиентская, 2",
            director_name="Сидоров Сергей Сергеевич",
            bank_name='АО "Банк клиента"',
            bik="044525111",
            settlement_account="40702810000000000002",
            correspondent_account="30101810000000000011",
        )
        self.carrier = Carrier.objects.create(
            name="Тестовый перевозчик",
            tax_id="7722000000",
            kpp="772201001",
            ogrn="1027700000002",
            address="Москва, ул. Транспортная, 3",
            director_name="Иванов Илья Ильич",
            bank_name='АО "Банк перевозчика"',
            bik="044525222",
            settlement_account="40702810000000000003",
            correspondent_account="30101810000000000022",
        )
        self.driver = Driver.objects.create(
            carrier=self.carrier,
            last_name="Иванов",
            first_name="Иван",
            phone="+7 900 000-00-01",
            license_number="TEST-DRIVER-1",
            license_categories="B, C, CE",
            license_expiry_date=date.today() + timedelta(days=365),
        )
        self.vehicle = Vehicle.objects.create(
            carrier=self.carrier,
            kind=Vehicle.Kind.TRACTOR,
            registration_number="Т001ЕСТ",
            make="Volvo",
            model="FH",
            capacity_kg=Decimal("20000"),
            volume_m3=Decimal("82"),
        )
        self.shipment = Shipment.objects.create(
            expeditor=self.company_profile,
            customer=self.customer,
            carrier=self.carrier,
            manager=self.user,
            cargo_name="Оборудование",
            pickup_city="Москва",
            pickup_date=date.today(),
            delivery_city="Казань",
            delivery_date=date.today() + timedelta(days=2),
            customer_price=Decimal("100000"),
            carrier_price=Decimal("75000"),
        )

    def create_second_expeditor(self):
        return CompanyProfile.objects.create(
            name='ООО "Второй экспедитор"',
            short_name='ООО "Второй экспедитор"',
            tax_id="7801000000",
            kpp="780101001",
            legal_address="Санкт-Петербург, ул. Логистическая, 5",
            default_vat_rate=CompanyProfile.VATRate.WITHOUT_VAT,
        )

    def create_shipment_for(self, expeditor, number="SECOND-001"):
        return Shipment.objects.create(
            number=number,
            expeditor=expeditor,
            customer=self.customer,
            manager=self.user,
            status=Shipment.Status.NEW,
            cargo_name="Груз второй компании",
            pickup_city="Санкт-Петербург",
            pickup_date=date.today(),
            delivery_city="Псков",
            delivery_date=date.today() + timedelta(days=1),
            customer_price=Decimal("300000"),
            carrier_price=Decimal("250000"),
        )

    def test_shipment_gets_number_and_margin(self):
        self.assertTrue(self.shipment.number.startswith("EXP-"))
        self.assertEqual(self.shipment.margin, Decimal("25000"))

    def test_dashboard_requires_login(self):
        response = self.client.get(reverse("dashboard"))
        self.assertRedirects(response, f"{reverse('login')}?next=/")

    def test_dashboard_and_lists_render(self):
        self.client.force_login(self.user)
        for url_name in (
            "dashboard", "shipment-list", "customer-list", "carrier-list",
            "driver-list", "vehicle-list", "reports", "shipment-document-list",
            "document-batch-list",
            "expeditor-list", "contract-list", "organization-list",
            "transportation-list", "order-list", "chat",
            "planner",
        ):
            with self.subTest(url_name=url_name):
                response = self.client.get(reverse(url_name))
                self.assertEqual(response.status_code, 200)

    def test_financial_sections_require_finance_role(self):
        logistician = get_user_model().objects.create_user(
            username="logistician", password="test-password"
        )
        logistician.crm_profile.role = UserProfile.Role.LOGISTICIAN
        logistician.crm_profile.save(update_fields=["role", "updated_at"])
        self.client.force_login(logistician)

        dashboard = self.client.get(reverse("dashboard"))
        self.assertEqual(dashboard.status_code, 200)
        self.assertNotContains(dashboard, reverse("bank-statement-list"))
        self.assertNotContains(dashboard, reverse("debt-report"))
        self.assertNotContains(dashboard, reverse("profitability-report"))

        for url_name in (
            "reports",
            "debt-report",
            "profitability-report",
            "bank-statement-list",
            "document-batch-list",
        ):
            with self.subTest(url_name=url_name):
                response = self.client.get(reverse(url_name))
                self.assertEqual(response.status_code, 403)

        accountant = get_user_model().objects.create_user(
            username="accountant", password="test-password"
        )
        accountant.crm_profile.role = UserProfile.Role.ACCOUNTANT
        accountant.crm_profile.save(update_fields=["role", "updated_at"])
        self.client.force_login(accountant)
        self.assertEqual(self.client.get(reverse("debt-report")).status_code, 200)

    def test_transportation_create_uses_modal_and_role_aware_searches(self):
        self.client.force_login(self.user)

        transportation_list = self.client.get(reverse("transportation-list"))
        form_response = self.client.get(reverse("transportation-create"))

        self.assertContains(transportation_list, "data-transportation-modal")
        self.assertContains(transportation_list, "js/transportation-workspace.js")
        self.assertContains(form_response, 'data-smart-select="organization"')
        self.assertContains(form_response, 'data-required-role="client"')
        self.assertContains(form_response, 'data-role-source="id_executor_role"')
        self.assertContains(form_response, 'data-parent-source="id_actual_carrier"')
        self.assertContains(form_response, 'data-contract-create-base="/contracts/new/"')
        self.assertContains(form_response, "Создать договор")
        self.assertContains(form_response, f"ИНН {self.customer.tax_id}")
        form = TransportationDocumentForm()
        self.assertIn(self.customer.organization, form.fields["client"].queryset)
        self.assertNotIn(
            self.carrier.organization,
            form.fields["client"].queryset,
        )
        self.assertEqual(form.fields["executor_amount"].widget.attrs["min"], "0")

    def test_order_form_supports_multiple_route_operations_and_registry(self):
        self.client.force_login(self.user)
        create_page = self.client.get(reverse("order-create"))
        self.assertEqual(create_page.status_code, 200)
        self.assertContains(create_page, "Ставка и форма оплаты")
        self.assertContains(create_page, "Добавить погрузку")
        self.assertContains(create_page, "Добавить выгрузку")
        self.assertEqual(
            len(create_page.context["stop_formset"].forms), 2
        )

        response = self.client.post(
            reverse("order-create"),
            {
                "owner_company": self.company_profile.organization.pk,
                "client": self.customer.organization.pk,
                "manager": self.user.pk,
                "document_date": date.today().isoformat(),
                "rate": "150000.00",
                "currency": "RUB",
                "payment_form": TransportOrder.PaymentForm.BANK_WITH_VAT,
                "payment_term_days": "15",
                "cargo_name": "Пластиковая тара",
                "cargo_description": "33 паллеты",
                "weight_kg": "15000",
                "volume_m3": "82",
                "package_count": "33",
                "pallet_count": "33",
                "package_type": "",
                "temperature_regime": "",
                "vehicle_requirements": "Тент",
                "special_requirements": "",
                "notes": "",
                "route_stops-TOTAL_FORMS": "3",
                "route_stops-INITIAL_FORMS": "0",
                "route_stops-MIN_NUM_FORMS": "0",
                "route_stops-MAX_NUM_FORMS": "1000",
                "route_stops-0-sequence": "1",
                "route_stops-0-kind": TransportOrderStop.Kind.PICKUP,
                "route_stops-0-city": "Санкт-Петербург",
                "route_stops-0-address": "Склад 1",
                "route_stops-0-planned_date": date.today().isoformat(),
                "route_stops-1-sequence": "2",
                "route_stops-1-kind": TransportOrderStop.Kind.PICKUP,
                "route_stops-1-city": "Тверь",
                "route_stops-1-address": "Склад 2",
                "route_stops-1-planned_date": (
                    date.today() + timedelta(days=1)
                ).isoformat(),
                "route_stops-2-sequence": "3",
                "route_stops-2-kind": TransportOrderStop.Kind.DELIVERY,
                "route_stops-2-city": "Москва",
                "route_stops-2-address": "Склад 3",
                "route_stops-2-planned_date": (
                    date.today() + timedelta(days=2)
                ).isoformat(),
                "action": "save",
            },
        )

        self.assertRedirects(response, reverse("order-list"))
        order = TransportOrder.objects.get(cargo_name="Пластиковая тара")
        self.assertRegex(order.number, r"^ЗК-\d{4}-\d{5}$")
        self.assertEqual(order.stops.count(), 3)
        self.assertEqual(order.route, "Санкт-Петербург → Тверь → Москва")
        registry = self.client.get(reverse("order-list"))
        self.assertContains(registry, order.number)
        self.assertContains(registry, "Пластиковая тара")
        self.assertContains(registry, "Ожидают назначения")
        edit_page = self.client.get(reverse("order-update", args=[order.pk]))
        self.assertEqual(edit_page.status_code, 200)
        self.assertContains(edit_page, f"Заказ {order.number}")
        self.assertContains(edit_page, "Пластиковая тара")
        self.assertContains(edit_page, "По клиенту нет действующего договора")
        self.assertContains(edit_page, "Создать договор")
        self.assertContains(edit_page, f"customer={self.customer.pk}")

    def test_order_assignment_creates_linked_trip_with_full_route(self):
        order = TransportOrder.objects.create(
            owner_company=self.company_profile.organization,
            client=self.customer.organization,
            manager=self.user,
            rate=Decimal("90000.00"),
            payment_form=TransportOrder.PaymentForm.BANK_WITHOUT_VAT,
            payment_term_days=10,
            cargo_name="Сборный груз",
            weight_kg=Decimal("5000"),
            planned_start_date=date.today(),
            planned_end_date=date.today() + timedelta(days=2),
        )
        for sequence, kind, city, time_from, time_to in (
            (1, TransportOrderStop.Kind.PICKUP, "Псков", time(8, 30), time(12, 0)),
            (2, TransportOrderStop.Kind.DELIVERY, "Тверь", time(14, 15), time(15, 45)),
            (3, TransportOrderStop.Kind.DELIVERY, "Москва", time(9, 0), time(11, 30)),
        ):
            TransportOrderStop.objects.create(
                order=order,
                sequence=sequence,
                kind=kind,
                city=city,
                planned_date=date.today() + timedelta(days=sequence - 1),
                planned_time_from=time_from,
                planned_time_to=time_to,
            )
        self.client.force_login(self.user)

        response = self.client.post(reverse("order-assign", args=[order.pk]))

        order.refresh_from_db()
        self.assertRedirects(
            response,
            reverse("transportation-update", args=[order.transportation_id]),
        )
        self.assertEqual(order.status, TransportOrder.Status.ASSIGNED)
        transportation = order.transportation
        self.assertEqual(transportation.customer_amount, Decimal("90000.00"))
        self.assertEqual(transportation.stops.count(), 3)
        self.assertEqual(
            list(transportation.stops.values_list("city", flat=True)),
            ["Псков", "Тверь", "Москва"],
        )
        first_stop = transportation.stops.order_by("sequence").first()
        self.assertEqual(
            timezone.localtime(first_stop.planned_from).time().replace(
                second=0, microsecond=0
            ),
            time(8, 30),
        )
        self.assertEqual(
            timezone.localtime(first_stop.planned_to).time().replace(
                second=0, microsecond=0
            ),
            time(12, 0),
        )
        self.assertTrue(
            transportation.parties.filter(
                organization=self.customer.organization,
                role=TransportationParty.Role.CLIENT,
                is_active=True,
            ).exists()
        )
        detail = self.client.get(
            reverse("transportation-detail", args=[transportation.pk])
        )
        self.assertContains(detail, f"заказ {order.number}")

        delete_page = self.client.get(reverse("order-delete", args=[order.pk]))
        self.assertContains(delete_page, "Удаление запрещено")
        self.assertContains(delete_page, "Созданный рейс")
        self.assertContains(delete_page, transportation.get_absolute_url())
        delete_response = self.client.post(
            reverse("order-delete", args=[order.pk])
        )
        self.assertEqual(delete_response.status_code, 409)
        self.assertTrue(TransportOrder.objects.filter(pk=order.pk).exists())

        trip_delete_url = reverse(
            "transportation-delete", args=[transportation.pk]
        )
        trip_delete_page = self.client.get(trip_delete_url)
        self.assertNotContains(trip_delete_page, "Удаление запрещено")
        self.assertContains(trip_delete_page, "будет возвращён в статус «Новый»")
        trip_delete_response = self.client.post(trip_delete_url)
        self.assertRedirects(trip_delete_response, reverse("transportation-list"))
        self.assertFalse(
            Transportation.objects.filter(pk=transportation.pk).exists()
        )
        order.refresh_from_db()
        self.assertEqual(order.status, TransportOrder.Status.NEW)
        self.assertIsNone(order.transportation_id)
        self.assertIsNone(order.assigned_at)
        self.assertIsNone(order.assigned_by_id)

    def test_quick_organization_create_adds_required_role_and_legacy_record(self):
        self.client.force_login(self.user)
        url = f"{reverse('quick-organization-create')}?role=client"

        response = self.client.post(
            url,
            {
                "kind": Organization.Kind.LEGAL_ENTITY,
                "name": 'ООО "Быстрый клиент"',
                "short_name": "Быстрый клиент",
                "tax_id": "7812345678",
                "kpp": "781201001",
                "roles": [OrganizationRole.Role.CLIENT],
                "is_active": "on",
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        organization = Organization.objects.get(pk=payload["item"]["id"])
        self.assertIn(OrganizationRole.Role.CLIENT, payload["item"]["roles"])
        self.assertTrue(
            organization.roles.filter(
                role=OrganizationRole.Role.CLIENT,
                is_active=True,
            ).exists()
        )
        self.assertTrue(Customer.objects.filter(organization=organization).exists())

    def test_quick_organization_create_rejects_duplicate_inn(self):
        self.client.force_login(self.user)
        response = self.client.post(
            f"{reverse('quick-organization-create')}?role=client",
            {
                "kind": Organization.Kind.LEGAL_ENTITY,
                "name": "Дубль",
                "tax_id": self.customer.tax_id,
                "roles": [OrganizationRole.Role.CLIENT],
                "is_active": "on",
            },
        )

        self.assertEqual(response.status_code, 422)
        self.assertContains(
            response,
            "Контрагент с таким ИНН уже существует",
            status_code=422,
        )

    def test_quick_driver_and_vehicle_create_for_selected_carrier(self):
        self.client.force_login(self.user)
        organization = self.carrier.organization
        driver_response = self.client.post(
            f"{reverse('quick-driver-create')}?organization={organization.pk}",
            {
                "last_name": "Петров",
                "first_name": "Пётр",
                "phone": "+7 900 111-22-33",
                "license_number": "QUICK-DRIVER",
                "license_categories": "C, CE",
                "license_expiry_date": "31.08.2028",
                "is_active": "on",
            },
        )
        vehicle_response = self.client.post(
            f"{reverse('quick-vehicle-create')}?organization={organization.pk}&resource_kind=vehicle",
            {
                "kind": Vehicle.Kind.TRACTOR,
                "registration_number": "А999АА77",
                "make": "КАМАЗ",
                "model": "54901",
                "capacity_kg": "20000",
                "volume_m3": "82",
                "pallet_capacity": "33",
                "is_active": "on",
            },
        )

        self.assertEqual(driver_response.status_code, 200)
        self.assertEqual(vehicle_response.status_code, 200)
        self.assertTrue(
            Driver.objects.filter(
                carrier=self.carrier,
                license_number="QUICK-DRIVER",
            ).exists()
        )
        self.assertTrue(
            Vehicle.objects.filter(
                carrier=self.carrier,
                registration_number="А999АА77",
            ).exists()
        )

    def test_dynamic_filters_are_enabled_across_list_pages_and_reports(self):
        self.client.force_login(self.user)
        for url_name in (
            "dashboard", "shipment-list", "customer-list", "carrier-list",
            "driver-list", "vehicle-list", "reports", "shipment-document-list",
            "expeditor-list", "contract-list", "organization-list",
            "transportation-list",
        ):
            with self.subTest(url_name=url_name):
                response = self.client.get(reverse(url_name))
                self.assertContains(response, "js/dynamic-filters.js")
                self.assertContains(response, "js/smart-selects.js")
        self.assertContains(
            self.client.get(reverse("dashboard")),
            "data-dynamic-filter",
        )

    def test_register_tables_include_shared_sorting_script(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("vehicle-list"))
        self.assertContains(response, "js/table-sort.js")
        self.assertContains(response, "v=20260906-money-sort")

    def test_expeditors_are_removed_from_sidebar_menu(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"))
        self.assertNotContains(response, 'href="/expeditors/"')
        self.assertContains(response, 'href="/organizations/"')

    def test_delete_buttons_are_available_for_accounting_entities(self):
        self.client.force_login(self.user)
        cases = (
            (reverse("shipment-detail", args=[self.shipment.pk]), reverse("shipment-delete", args=[self.shipment.pk])),
            (self.shipment.transportation.get_absolute_url(), reverse("transportation-delete", args=[self.shipment.transportation.pk])),
            (self.customer.organization.get_absolute_url(), reverse("organization-delete", args=[self.customer.organization.pk])),
            (self.carrier.get_absolute_url(), reverse("organization-delete", args=[self.carrier.organization.pk])),
            (self.driver.get_absolute_url(), reverse("driver-delete", args=[self.driver.pk])),
            (self.vehicle.get_absolute_url(), reverse("vehicle-delete", args=[self.vehicle.pk])),
        )
        for page_url, delete_url in cases:
            with self.subTest(page_url=page_url):
                self.assertContains(self.client.get(page_url), delete_url)

    def test_linked_driver_and_vehicle_cannot_be_deleted(self):
        self.client.force_login(self.user)
        Shipment.objects.filter(pk=self.shipment.pk).update(
            driver=self.driver,
            vehicle=self.vehicle,
        )
        for model, url_name in (
            (self.driver, "driver-delete"),
            (self.vehicle, "vehicle-delete"),
        ):
            with self.subTest(url_name=url_name):
                url = reverse(url_name, args=[model.pk])
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "Удаление запрещено")
                self.assertContains(response, "Заявки")
                post_response = self.client.post(url)
                self.assertEqual(post_response.status_code, 409)
                self.assertTrue(type(model).objects.filter(pk=model.pk).exists())

    def test_unlinked_driver_and_vehicle_can_be_deleted(self):
        self.client.force_login(self.user)
        driver = Driver.objects.create(
            carrier=self.carrier,
            last_name="Свободный",
            first_name="Водитель",
            phone="+7 900 111-11-11",
            license_number="FREE-DRIVER",
            license_categories="C",
            license_expiry_date=date.today() + timedelta(days=365),
        )
        vehicle = Vehicle.objects.create(
            carrier=self.carrier,
            kind=Vehicle.Kind.TRACTOR,
            registration_number="С002ВОБ",
            make="MAN",
        )
        for model, url_name, list_name in (
            (driver, "driver-delete", "driver-list"),
            (vehicle, "vehicle-delete", "vehicle-list"),
        ):
            with self.subTest(url_name=url_name):
                response = self.client.post(reverse(url_name, args=[model.pk]))
                self.assertRedirects(response, reverse(list_name))
                self.assertFalse(type(model).objects.filter(pk=model.pk).exists())

    def test_linked_organization_cannot_be_deleted(self):
        self.client.force_login(self.user)
        organization = self.customer.organization
        url = reverse("organization-delete", args=[organization.pk])
        response = self.client.get(url)
        self.assertContains(response, "Удаление запрещено")
        self.assertContains(response, "Заявки / рейсы")
        self.assertEqual(self.client.post(url).status_code, 409)
        self.assertTrue(Organization.objects.filter(pk=organization.pk).exists())

    def test_unlinked_organization_and_technical_card_can_be_deleted(self):
        self.client.force_login(self.user)
        organization = Organization.objects.create(
            name="Свободный контрагент",
            tax_id="7800000088",
        )
        legacy_customer = Customer.objects.create(
            name=organization.name,
            tax_id=organization.tax_id,
            organization=organization,
        )
        response = self.client.post(
            reverse("organization-delete", args=[organization.pk])
        )
        self.assertRedirects(response, reverse("organization-list"))
        self.assertFalse(Organization.objects.filter(pk=organization.pk).exists())
        self.assertFalse(Customer.objects.filter(pk=legacy_customer.pk).exists())

    def test_draft_shipment_and_mirrored_trip_are_deleted_together(self):
        self.client.force_login(self.user)
        shipment_id = self.shipment.pk
        transportation_id = self.shipment.transportation.pk
        response = self.client.post(
            reverse("shipment-delete", args=[shipment_id])
        )
        self.assertRedirects(response, reverse("shipment-list"))
        self.assertFalse(Shipment.objects.filter(pk=shipment_id).exists())
        self.assertFalse(Transportation.objects.filter(pk=transportation_id).exists())

    def test_posted_transportation_cannot_be_deleted(self):
        self.client.force_login(self.user)
        transportation = self.shipment.transportation
        Transportation.objects.filter(pk=transportation.pk).update(
            posting_status=Transportation.PostingStatus.POSTED
        )
        url = reverse("transportation-delete", args=[transportation.pk])
        response = self.client.get(url)
        self.assertContains(response, "Удаление запрещено")
        self.assertContains(response, "Состояние заявки / рейса")
        self.assertEqual(self.client.post(url).status_code, 409)
        self.assertTrue(Transportation.objects.filter(pk=transportation.pk).exists())

    def test_legacy_shipment_is_mirrored_as_transportation(self):
        transportation = self.shipment.transportation
        self.assertEqual(transportation.number, self.shipment.number)
        self.assertEqual(transportation.owner_company, self.company_profile.organization)
        self.assertEqual(transportation.route, "Москва → Казань")
        self.assertEqual(transportation.stops.count(), 2)
        self.assertTrue(
            transportation.parties.filter(
                role=TransportationParty.Role.CLIENT,
                organization=self.customer.organization,
            ).exists()
        )
        link = transportation.execution_links.get(sequence=1)
        self.assertEqual(link.contractor_role, TransportationLink.ContractorRole.CARRIER)
        self.assertEqual(link.contractor_party.organization, self.carrier.organization)
        assignment = transportation.vehicle_assignments.get(is_active=True)
        self.assertEqual(assignment.actual_carrier, self.carrier.organization)

    def test_organizations_with_same_inn_receive_multiple_roles(self):
        shared_customer = Customer.objects.create(
            name="Многофункциональная компания",
            tax_id="7800000099",
        )
        shared_carrier = Carrier.objects.create(
            name="Многофункциональная компания",
            tax_id="7800000099",
        )
        self.assertEqual(shared_customer.organization, shared_carrier.organization)
        roles = set(
            shared_customer.organization.roles.filter(is_active=True).values_list(
                "role", flat=True
            )
        )
        self.assertEqual(
            roles,
            {OrganizationRole.Role.CLIENT, OrganizationRole.Role.CARRIER},
        )

    def test_unified_organization_form_creates_legacy_directory_records(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("organization-create"),
            {
                "name": 'ООО "Новая роль"',
                "short_name": 'ООО "Новая роль"',
                "tax_id": "7812000099",
                "kpp": "781201001",
                "ogrn": "1027800000099",
                "legal_address": "Санкт-Петербург, тестовый адрес",
                "director_name": "Ролевой Роман Романович",
                "contact_name": "Роман",
                "phone": "+7 900 000-00-99",
                "email": "roles@example.com",
                "bank_name": "Тест Банк",
                "bik": "044030099",
                "settlement_account": "40702810000000000099",
                "correspondent_account": "30101810000000000099",
                "verification_status": Organization.VerificationStatus.VERIFIED,
                "roles": [
                    OrganizationRole.Role.CLIENT,
                    OrganizationRole.Role.FORWARDER,
                    OrganizationRole.Role.CARRIER,
                ],
                "is_active": "on",
            },
        )
        organization = Organization.objects.get(tax_id="7812000099")
        self.assertRedirects(response, organization.get_absolute_url())
        self.assertTrue(organization.legacy_customers.exists())
        self.assertTrue(organization.legacy_carriers.exists())
        self.assertEqual(organization.roles.filter(is_active=True).count(), 3)

    def test_organization_workspace_saves_primary_bank_and_contact_registers(self):
        self.client.force_login(self.user)
        organization = self.customer.organization
        response = self.client.post(
            reverse("organization-update", args=[organization.pk]),
            {
                "kind": Organization.Kind.LEGAL_ENTITY,
                "name": organization.name,
                "short_name": "Клиент в программе",
                "registration_country": "РОССИЯ",
                "tax_id": organization.tax_id,
                "kpp": organization.kpp,
                "roles": [OrganizationRole.Role.CLIENT],
                "is_active": "on",
                "verification_status": Organization.VerificationStatus.NOT_CHECKED,
                "bank_accounts-TOTAL_FORMS": "1",
                "bank_accounts-INITIAL_FORMS": "0",
                "bank_accounts-MIN_NUM_FORMS": "0",
                "bank_accounts-MAX_NUM_FORMS": "1000",
                "bank_accounts-0-account_number": "40702810000000000999",
                "bank_accounts-0-bank_name": 'АО "Новый банк"',
                "bank_accounts-0-bik": "044525999",
                "bank_accounts-0-correspondent_account": "30101810000000000999",
                "bank_accounts-0-currency": "RUB",
                "bank_accounts-0-is_primary": "on",
                "bank_accounts-0-is_active": "on",
                "contact_people-TOTAL_FORMS": "1",
                "contact_people-INITIAL_FORMS": "0",
                "contact_people-MIN_NUM_FORMS": "0",
                "contact_people-MAX_NUM_FORMS": "1000",
                "contact_people-0-full_name": "Кузнецова Елена Игоревна",
                "contact_people-0-position": "Главный бухгалтер",
                "contact_people-0-phone": "+7 900 111-22-33",
                "contact_people-0-email": "finance@example.com",
                "contact_people-0-is_primary": "on",
                "contact_people-0-is_active": "on",
                "action": "save",
            },
        )
        self.assertRedirects(response, organization.get_absolute_url())
        bank = OrganizationBankAccount.objects.get(organization=organization)
        contact = OrganizationContact.objects.get(organization=organization)
        self.assertTrue(bank.is_primary)
        self.assertTrue(contact.is_primary)
        organization.refresh_from_db()
        self.assertEqual(organization.settlement_account, bank.account_number)
        self.assertEqual(organization.contact_name, contact.full_name)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.settlement_account, bank.account_number)

    def test_organization_workspace_contains_1c_sections_and_balances(self):
        self.client.force_login(self.user)
        group = OrganizationGroup.objects.create(name="Покупатели")
        organization = self.customer.organization
        organization.group = group
        organization.save(update_fields=["group", "updated_at"])
        response = self.client.get(organization.get_absolute_url())
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Записать и закрыть")
        self.assertContains(response, "Банковские счета")
        self.assertContains(response, "Контактные лица")
        self.assertContains(response, "Взаиморасчёты")
        self.assertEqual(response.context["receivable_total"], Decimal("100000"))

        carrier_response = self.client.get(self.carrier.organization.get_absolute_url())
        self.assertEqual(carrier_response.context["payable_total"], Decimal("75000"))

    def test_organization_save_and_close_returns_to_directory(self):
        self.client.force_login(self.user)
        organization = self.customer.organization
        response = self.client.post(
            reverse("organization-update", args=[organization.pk]),
            {
                "kind": Organization.Kind.LEGAL_ENTITY,
                "name": organization.name,
                "registration_country": "РОССИЯ",
                "tax_id": organization.tax_id,
                "roles": [OrganizationRole.Role.CLIENT],
                "is_active": "on",
                "verification_status": Organization.VerificationStatus.NOT_CHECKED,
                "action": "save_close",
            },
        )
        self.assertRedirects(response, reverse("organization-list"))

    def test_counterparty_roles_use_standard_checkboxes_without_payer(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("organization-update", args=[self.customer.organization.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Плательщик")
        self.assertContains(response, 'name="roles"', count=5)
        self.assertEqual(
            response.context["form"]["roles"].value(),
            [OrganizationRole.Role.CLIENT],
        )
        self.assertContains(response, 'id="id_roles_0" checked')
        self.assertContains(response, 'class="checkbox uk-checkbox"')
        self.assertEqual(
            [value for value, _label in OrganizationRole.Role.choices],
            ["client", "forwarder", "carrier", "shipper", "consignee"],
        )

    def test_counterparty_roles_remain_checked_after_update(self):
        self.client.force_login(self.user)
        organization = self.customer.organization
        response = self.client.post(
            reverse("organization-update", args=[organization.pk]),
            {
                "kind": Organization.Kind.LEGAL_ENTITY,
                "name": organization.name,
                "registration_country": "РОССИЯ",
                "tax_id": organization.tax_id,
                "roles": [
                    OrganizationRole.Role.CLIENT,
                    OrganizationRole.Role.FORWARDER,
                ],
                "is_active": "on",
                "verification_status": Organization.VerificationStatus.NOT_CHECKED,
                "action": "save",
            },
        )
        self.assertRedirects(response, organization.get_absolute_url())

        edit_response = self.client.get(organization.get_absolute_url())
        self.assertCountEqual(
            edit_response.context["form"]["roles"].value(),
            [OrganizationRole.Role.CLIENT, OrganizationRole.Role.FORWARDER],
        )
        self.assertContains(edit_response, 'id="id_roles_0" checked')
        self.assertContains(edit_response, 'id="id_roles_1" checked')

    def test_duplicate_counterparty_inn_is_rejected_with_existing_card_link(self):
        self.client.force_login(self.user)
        organization = self.customer.organization
        response = self.client.post(
            reverse("organization-create"),
            {
                "kind": Organization.Kind.LEGAL_ENTITY,
                "name": "Контрагент-дубль",
                "registration_country": "РОССИЯ",
                "tax_id": f"{organization.tax_id[:4]} {organization.tax_id[4:]}",
                "roles": [OrganizationRole.Role.CLIENT],
                "is_active": "on",
                "verification_status": Organization.VerificationStatus.NOT_CHECKED,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertFormError(
            response.context["form"],
            "tax_id",
            "Контрагент с таким ИНН уже существует. Откройте его карточку.",
        )
        self.assertContains(response, str(organization))
        self.assertContains(response, organization.get_absolute_url())
        self.assertFalse(Organization.objects.filter(name="Контрагент-дубль").exists())

    def test_counterparty_inn_live_check_returns_existing_card(self):
        self.client.force_login(self.user)
        organization = self.customer.organization
        response = self.client.get(
            reverse("organization-check-inn"),
            {"inn": organization.tax_id},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["exists"])
        self.assertEqual(
            response.json()["organization"]["url"], organization.get_absolute_url()
        )

        edit_response = self.client.get(
            reverse("organization-check-inn"),
            {"inn": organization.tax_id, "exclude": organization.pk},
        )
        self.assertFalse(edit_response.json()["exists"])

    def test_counterparty_inn_is_unique_at_database_level(self):
        organization = self.customer.organization
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Organization.objects.create(
                    name="Контрагент-дубль",
                    tax_id=f"{organization.tax_id[:4]}-{organization.tax_id[4:]}",
                )

    def test_chain_form_separates_forwarder_and_actual_carrier(self):
        forwarder = Carrier.objects.create(
            name='ООО "Привлечённый экспедитор"',
            tax_id="7722000098",
        )
        OrganizationRole.objects.update_or_create(
            organization=forwarder.organization,
            role=OrganizationRole.Role.FORWARDER,
            defaults={"is_active": True},
        )
        transportation = self.shipment.transportation
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("transportation-chain-update", args=[transportation.pk]),
            {
                "executor": forwarder.organization.pk,
                "executor_role": TransportationLink.ContractorRole.FORWARDER,
                "contract": "",
                "instruction_number": "ПЭ-001",
                "instruction_status": "Принято",
                "actual_carrier": self.carrier.organization.pk,
                "driver": self.driver.pk,
                "vehicle": self.vehicle.pk,
                "trailer": "",
                "trailer_registration_number": "В456ВВ198",
            },
        )
        self.assertRedirects(response, transportation.get_absolute_url())
        transportation.refresh_from_db()
        self.assertFalse(transportation.legacy_chain_sync)
        links = list(transportation.execution_links.order_by("sequence"))
        self.assertEqual(len(links), 2)
        self.assertEqual(
            links[0].contractor_role, TransportationLink.ContractorRole.FORWARDER
        )
        self.assertEqual(
            links[1].contractor_party.organization, self.carrier.organization
        )
        assignment = transportation.vehicle_assignments.get(is_active=True)
        self.assertEqual(assignment.actual_carrier, self.carrier.organization)
        self.assertEqual(assignment.driver, self.driver)

    def test_direct_carrier_form_rejects_different_actual_carrier(self):
        other = Carrier.objects.create(name="Другой перевозчик", tax_id="7700000097")
        form = TransportationChainForm(
            data={
                "executor": self.carrier.organization.pk,
                "executor_role": TransportationLink.ContractorRole.CARRIER,
                "actual_carrier": other.organization.pk,
            },
            transportation=self.shipment.transportation,
        )
        self.assertFalse(form.is_valid())
        self.assertIn("actual_carrier", form.errors)

    def test_transportation_posting_requires_executor_contract(self):
        self.client.force_login(self.user)
        transportation = self.shipment.transportation
        response = self.client.post(
            reverse("transportation-post", args=[transportation.pk])
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            response["Location"],
            reverse("transportation-update", args=[transportation.pk]),
        )
        transportation.refresh_from_db()
        self.assertEqual(
            transportation.posting_status, Transportation.PostingStatus.DRAFT
        )
        follow_up = self.client.get(response.url)
        self.assertContains(follow_up, "Создайте и выберите договор")

    def test_ezz_can_be_prepared_before_driver_and_vehicle_are_known(self):
        transportation = self.shipment.transportation
        transportation.weight_kg = Decimal("100")
        transportation.save(update_fields=["weight_kg", "updated_at"])
        transportation.stops.update(address="Склад")
        old_assignment = transportation.active_vehicle_assignment()
        vehicle_number = (
            old_assignment.vehicle.registration_number
            if old_assignment and old_assignment.vehicle_id
            else ""
        )
        transportation.vehicle_assignments.update(is_active=False)

        self.assertEqual(
            epd_validation_errors(
                transportation,
                kind=TransportationElectronicDocument.Kind.EZZ,
            ),
            [],
        )
        self.assertTrue(
            any(
                "Не назначены водитель и транспорт" in message
                for message in epd_validation_errors(
                    transportation,
                    kind=TransportationElectronicDocument.Kind.ETRN,
                )
            )
        )

        documents = prepare_documents(transportation, self.user)
        ezz = next(
            document
            for document in documents
            if document.kind == TransportationElectronicDocument.Kind.EZZ
        )
        self.assertNotIn("<driver>", ezz.operator_xml)
        if vehicle_number:
            self.assertNotIn(vehicle_number, ezz.operator_xml)
        self.assertFalse(
            any(
                document.kind == TransportationElectronicDocument.Kind.ETRN
                for document in documents
            )
        )

        self.client.force_login(self.user)
        response = self.client.get(
            reverse(
                "transportation-epd-download",
                args=[transportation.pk, ezz.pk],
            )
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"<\xd4\xe0\xe9\xeb", response.content)
        self.assertNotIn(b"LogisticsOrderRequestSenderTitle", response.content)

        response = self.client.get(
            reverse(
                "transportation-epd-download",
                args=[transportation.pk, ezz.pk],
            )
            + "?format=userdata"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"LogisticsOrderRequestSenderTitle", response.content)
        self.assertNotIn(b"<driver>", response.content)

    @override_settings(
        KONTUR_DIADOC_API_TOKEN="test-token",
        KONTUR_DIADOC_API_URL="https://diadoc.test",
        KONTUR_DIADOC_API_TIMEOUT=3,
    )
    @patch("crm.kontur.urlopen")
    def test_kontur_generate_title_uses_userdata_and_decodes_response(
        self, mock_urlopen
    ):
        from .kontur import generate_ezz_title_xml

        owner = self.company_profile.organization
        owner.edo_id = "2BM-test-box"
        owner.save(update_fields=["edo_id", "updated_at"])
        final_xml = '<?xml version="1.0" encoding="windows-1251"?>\n<Файл />'
        mock_urlopen.return_value.__enter__.return_value.read.return_value = (
            final_xml.encode("cp1251")
        )

        result = generate_ezz_title_xml(
            self.shipment.transportation,
            '<LogisticsOrderRequestSenderTitle Number="TEST" />',
        )

        self.assertIn("<Файл", result)
        request = mock_urlopen.call_args.args[0]
        self.assertIn("/GenerateTitleXml?", request.full_url)
        self.assertIn("documentTypeNamedId=LogisticsOrderRequest", request.full_url)
        self.assertEqual(request.get_header("Authorization"), "Bearer test-token")
        self.assertEqual(
            request.data,
            b'<LogisticsOrderRequestSenderTitle Number="TEST" />',
        )

    @override_settings(
        KONTUR_DIADOC_API_TOKEN="test-token",
        KONTUR_DIADOC_API_URL="https://diadoc.test",
    )
    @patch("crm.views.generate_ezz_title_xml")
    def test_ezz_download_caches_final_kontur_xml(self, mock_generate):
        transportation = self.shipment.transportation
        owner = transportation.owner_company
        owner.edo_id = "2BM-test-box"
        owner.save(update_fields=["edo_id", "updated_at"])
        document = TransportationElectronicDocument.objects.create(
            transportation=transportation,
            kind=TransportationElectronicDocument.Kind.EZZ,
            status=TransportationElectronicDocument.Status.READY,
            operator_xml='<LogisticsOrderRequestSenderTitle Number="TEST" />',
        )
        final_xml = '<?xml version="1.0" encoding="windows-1251"?>\n<Файл />'
        mock_generate.return_value = final_xml
        self.client.force_login(self.user)

        response = self.client.get(document.get_absolute_url())

        self.assertEqual(response.status_code, 200)
        self.assertIn("ON_ZAKZVGO-", response["Content-Disposition"])
        self.assertEqual(response.content, final_xml.encode("cp1251"))
        document.refresh_from_db()
        self.assertEqual(document.generated_xml, final_xml)
        self.assertEqual(
            document.provider,
            TransportationElectronicDocument.Provider.KONTUR,
        )
        mock_generate.assert_called_once()

        response = self.client.get(document.get_absolute_url())
        self.assertEqual(response.content, final_xml.encode("cp1251"))
        mock_generate.assert_called_once()

    def test_contract_can_be_created_from_transportation_chain(self):
        self.client.force_login(self.user)
        transportation = self.shipment.transportation
        chain_page = self.client.get(
            reverse("transportation-chain-update", args=[transportation.pk])
        )
        self.assertContains(chain_page, "Создать договор")

        contract_url = (
            f"{reverse('contract-create')}?kind={Contract.Kind.CARRIER_TRANSPORT}"
            f"&return_transportation={transportation.pk}"
            f"&owner_organization={transportation.owner_company_id}"
            f"&counterparty_organization={self.carrier.organization_id}"
        )
        contract_page = self.client.get(contract_url)
        self.assertEqual(
            contract_page.context["form"].initial["expeditor"],
            self.company_profile.pk,
        )
        self.assertEqual(
            contract_page.context["form"].initial["carrier"],
            self.carrier.pk,
        )
        response = self.client.post(
            contract_url,
            {
                "kind": Contract.Kind.CARRIER_TRANSPORT,
                "expeditor": self.company_profile.pk,
                "carrier": self.carrier.pk,
                "number": "ДГ-ИЗ-РЕЙС-001",
                "contract_date": date.today().isoformat(),
                "city": "Москва",
                "valid_until": "",
                "status": Contract.Status.DRAFT,
                "expeditor_representative": self.company_profile.director_name,
                "expeditor_authority_basis": "Устава",
                "counterparty_representative": self.carrier.director_name,
                "counterparty_authority_basis": "Устава",
                "payment_term_days": "5",
                "vat_rate": "",
                "notes": "Создан из карточки рейса",
            },
        )
        contract = Contract.objects.get(number="ДГ-ИЗ-РЕЙС-001")
        expected_chain_url = (
            f"{reverse('transportation-chain-update', args=[transportation.pk])}"
            f"?contract={contract.pk}"
        )
        self.assertRedirects(response, expected_chain_url)
        chain_after = self.client.get(response.url)
        self.assertEqual(chain_after.context["form"].initial["contract"], contract.pk)

    def test_new_transportation_and_organization_cards_render(self):
        self.client.force_login(self.user)
        transportation = self.shipment.transportation
        pages = (
            transportation.get_absolute_url(),
            reverse("transportation-chain-update", args=[transportation.pk]),
            self.customer.organization.get_absolute_url(),
            reverse("organization-update", args=[self.customer.organization.pk]),
        )
        for url in pages:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
        detail = self.client.get(transportation.get_absolute_url())
        self.assertContains(detail, "Цепочка перевозки")
        self.assertContains(detail, "Фактический перевозчик")
        self.assertContains(detail, "НДС к уплате")
        self.assertContains(detail, "Прибыль без НДС")
        self.assertContains(detail, "История")
        self.assertContains(detail, self.shipment.number)
        self.assertContains(detail, "transportation-card")
        self.assertContains(detail, "onec-titlebar")
        self.assertContains(detail, "onec-commandbar")
        self.assertContains(detail, "transportation-detail-tabs")
        self.assertContains(detail, "Назначить исполнение")
        self.assertContains(detail, "Назначить водителя и ТС")
        self.assertContains(detail, "Проведение документа")
        self.assertContains(detail, "До проведения нужно заполнить")

    def test_transportation_register_has_counterparty_style_scopes(self):
        self.client.force_login(self.user)
        register = self.client.get(reverse("transportation-list"))
        self.assertContains(register, "Реестр рейсов")
        self.assertContains(register, "Создать рейс")
        self.assertContains(register, "scope=active")
        self.assertEqual(register.context["transportation_count"], 1)
        self.assertEqual(register.context["active_count"], 1)
        active = self.client.get(reverse("transportation-list"), {"scope": "active"})
        self.assertEqual(active.context["page_obj"].paginator.count, 1)
        closed = self.client.get(reverse("transportation-list"), {"scope": "closed"})
        self.assertEqual(closed.context["page_obj"].paginator.count, 0)

    def test_transportation_status_flow_rejects_skipping_stages(self):
        transportation = self.shipment.transportation
        self.assertEqual(transportation.status, Transportation.Status.NEW)
        self.assertEqual(
            transportation.status_transition_issues(
                Transportation.Status.CLIENT_INSTRUCTION_RECEIVED
            ),
            [],
        )
        issues = transportation.status_transition_issues(
            Transportation.Status.VEHICLE_CONFIRMED
        )
        self.assertEqual(len(issues), 1)
        self.assertIn("Поручение клиента получено", issues[0])

    def test_posted_transportation_can_advance_only_to_next_stage(self):
        self.shipment.driver = self.driver
        self.shipment.vehicle = self.vehicle
        self.shipment.save()
        transportation = self.shipment.transportation
        assignment = transportation.active_vehicle_assignment()
        assignment.trailer_registration_number = "В456ВВ198"
        assignment.save(update_fields=["trailer_registration_number", "updated_at"])
        contract = Contract.objects.create(
            kind=Contract.Kind.CARRIER_TRANSPORT,
            expeditor=self.company_profile,
            carrier=self.carrier,
            contract_date=date.today(),
        )
        link = transportation.active_execution_link()
        link.contract = contract
        link.save(update_fields=["contract", "updated_at"])
        self.client.force_login(self.user)

        post_response = self.client.post(
            reverse("transportation-post", args=[transportation.pk])
        )
        self.assertRedirects(post_response, transportation.get_absolute_url())
        transportation.refresh_from_db()
        self.assertEqual(transportation.posting_status, Transportation.PostingStatus.POSTED)

        invalid = self.client.post(
            reverse("transportation-advance-status", args=[transportation.pk]),
            {"target_status": Transportation.Status.CLOSED},
        )
        self.assertRedirects(invalid, transportation.get_absolute_url())
        transportation.refresh_from_db()
        self.assertEqual(transportation.status, Transportation.Status.NEW)
        self.assertContains(
            self.client.get(transportation.get_absolute_url()),
            "Следующий этап: Поручение клиента получено",
        )

        transportation.status = Transportation.Status.VEHICLE_CONFIRMED
        transportation.save(update_fields=["status", "updated_at"])
        advance = self.client.post(
            reverse("transportation-advance-status", args=[transportation.pk]),
            {"target_status": Transportation.Status.LOADING},
        )
        self.assertRedirects(advance, transportation.get_absolute_url())
        transportation.refresh_from_db()
        self.assertEqual(transportation.status, Transportation.Status.LOADING)
        self.assertTrue(
            transportation.status_events.filter(
                source=TransportationStatusEvent.Source.MANUAL,
                new_status=Transportation.Status.LOADING,
            ).exists()
        )

    def test_transportation_close_requires_final_accounting_checks(self):
        self.shipment.driver = self.driver
        self.shipment.vehicle = self.vehicle
        self.shipment.save()
        transportation = self.shipment.transportation
        assignment = transportation.active_vehicle_assignment()
        assignment.trailer_registration_number = "В456ВВ198"
        assignment.save(update_fields=["trailer_registration_number", "updated_at"])
        contract = Contract.objects.create(
            kind=Contract.Kind.CARRIER_TRANSPORT,
            expeditor=self.company_profile,
            carrier=self.carrier,
            contract_date=date.today(),
        )
        link = transportation.active_execution_link()
        link.contract = contract
        link.save(update_fields=["contract", "updated_at"])
        self.client.force_login(self.user)

        self.client.post(reverse("transportation-post", args=[transportation.pk]))
        transportation.refresh_from_db()
        response = self.client.post(reverse("transportation-close", args=[transportation.pk]))
        self.assertRedirects(response, transportation.get_absolute_url())
        transportation.refresh_from_db()
        self.assertNotEqual(transportation.status, Transportation.Status.CLOSED)
        detail = self.client.get(transportation.get_absolute_url())
        self.assertContains(detail, "Закрытие рейса")
        self.assertContains(detail, "Груз должен быть доставлен")

        transportation.status = Transportation.Status.CUSTOMER_INVOICED
        transportation.save(update_fields=["status", "updated_at"])
        Payment.objects.create(
            transportation=transportation,
            direction=Payment.Direction.INCOME,
            amount=transportation.customer_amount,
            payment_date=date.today(),
            reference="ПП-CLIENT-001",
            created_by=self.user,
        )
        Payment.objects.create(
            transportation=transportation,
            direction=Payment.Direction.EXPENSE,
            amount=transportation.executor_amount,
            payment_date=date.today(),
            reference="ПП-CARRIER-001",
            created_by=self.user,
        )

        close_response = self.client.post(
            reverse("transportation-close", args=[transportation.pk])
        )
        self.assertRedirects(close_response, transportation.get_absolute_url())
        transportation.refresh_from_db()
        self.assertEqual(transportation.status, Transportation.Status.CLOSED)
        self.assertTrue(
            transportation.status_events.filter(
                new_status=Transportation.Status.CLOSED,
                comment="Рейс закрыт",
            ).exists()
        )

    def test_erp_transportation_document_posts_register_movements(self):
        self.client.force_login(self.user)
        document_date = date.today()
        delivery_date = document_date + timedelta(days=2)
        executor_contract = Contract.objects.create(
            kind=Contract.Kind.CARRIER_TRANSPORT,
            expeditor=self.company_profile,
            carrier=self.carrier,
            contract_date=document_date,
        )
        response = self.client.post(
            reverse("transportation-create"),
            {
                "owner_company": self.company_profile.organization.pk,
                "manager": self.user.pk,
                "document_date": document_date.isoformat(),
                "status": Transportation.Status.VEHICLE_CONFIRMED,
                "client": self.customer.organization.pk,
                "client_reference": "КЛ-2026-001",
                "customer_contract": "",
                "customer_amount": "60000.00",
                "customer_vat_rate": VATRate.objects.get(code="22").pk,
                "customer_payment_term_days": "10",
                "pickup_city": "Санкт-Петербург",
                "pickup_address": "Склад клиента",
                "pickup_date": document_date.isoformat(),
                "delivery_city": "Москва",
                "delivery_address": "Склад получателя",
                "delivery_date": delivery_date.isoformat(),
                "cargo_name": "Пластиковая тара",
                "cargo_description": "33 паллеты",
                "weight_kg": "15000.00",
                "volume_m3": "82.00",
                "package_count": "33",
                "pallet_count": "33",
                "package_type": "",
                "loading_method": "",
                "unloading_method": "",
                "temperature_regime": "",
                "vehicle_requirements": "Тент",
                "special_requirements": "",
                "executor": self.carrier.organization.pk,
                "executor_role": TransportationLink.ContractorRole.CARRIER,
                "executor_contract": executor_contract.pk,
                "executor_instruction_number": "ИСП-001",
                "executor_amount": "40000.00",
                "executor_vat_rate": VATRate.objects.get(code="without_vat").pk,
                "executor_payment_term_days": "5",
                "payment_due_basis": Transportation.PaymentDueBasis.DELIVERY_DATE,
                "actual_carrier": self.carrier.organization.pk,
                "driver": self.driver.pk,
                "vehicle": self.vehicle.pk,
                "trailer": "",
                "trailer_registration_number": "В456ВВ198",
                "currency": "RUB",
                "notes": "",
                "action": "post",
            },
        )
        transportation = Transportation.objects.get(
            client_reference="КЛ-2026-001"
        )
        document_event = transportation.status_events.get(source="document")
        self.assertEqual(document_event.comment, "Документ создан")
        self.assertEqual(
            document_event.changes["Груз"]["new"], "Пластиковая тара"
        )
        self.assertEqual(
            document_event.changes["Клиент"]["new"], str(self.customer.organization)
        )
        self.assertRedirects(response, transportation.get_absolute_url())
        self.assertEqual(
            transportation.posting_status, Transportation.PostingStatus.POSTED
        )
        self.assertRegex(transportation.number, r"^\d{2}/\d{2}-\d{4}$")
        self.assertEqual(transportation.margin, Decimal("20000.00"))
        self.assertEqual(transportation.customer_vat_amount, Decimal("10819.67"))
        self.assertEqual(transportation.executor_vat_amount, Decimal("0.00"))
        self.assertEqual(transportation.vat_payable, Decimal("10819.67"))
        self.assertEqual(transportation.profit, Decimal("9180.33"))
        self.assertEqual(transportation.profit_tax_amount, Decimal("2295.08"))
        self.assertEqual(transportation.net_profit, Decimal("6885.25"))
        self.assertEqual(
            transportation.customer_payment_due_date,
            delivery_date + timedelta(days=10),
        )
        self.assertEqual(
            transportation.executor_payment_due_date,
            delivery_date + timedelta(days=5),
        )
        self.assertEqual(transportation.charges.count(), 2)
        self.assertEqual(transportation.settlement_movements.count(), 2)
        self.assertEqual(transportation.instructions.count(), 2)
        self.assertEqual(transportation.receivable_balance, Decimal("60000.00"))
        self.assertEqual(transportation.payable_balance, Decimal("40000.00"))
        self.assertTrue(
            transportation.charges.filter(
                direction=TripCharge.Direction.REVENUE,
                amount=Decimal("60000.00"),
            ).exists()
        )
        self.assertTrue(
            transportation.settlement_movements.filter(
                side=SettlementMovement.Side.PAYABLE,
                kind=SettlementMovement.Kind.ACCRUAL,
            ).exists()
        )
        self.assertTrue(
            transportation.instructions.filter(
                kind=TransportationInstruction.Kind.CARRIER_APPLICATION
            ).exists()
        )

        repeat = self.client.post(
            reverse("transportation-post", args=[transportation.pk])
        )
        self.assertRedirects(repeat, transportation.get_absolute_url())
        self.assertEqual(transportation.charges.count(), 2)
        self.assertEqual(transportation.settlement_movements.count(), 2)

        payment_page = self.client.get(
            reverse("transportation-payment-create", args=[transportation.pk])
            + "?direction=income"
        )
        self.assertEqual(payment_page.status_code, 200)
        self.assertContains(payment_page, "Зарегистрировать платёж")
        payment_response = self.client.post(
            reverse("transportation-payment-create", args=[transportation.pk]),
            {
                "direction": Payment.Direction.INCOME,
                "amount": "15000.00",
                "payment_date": document_date.isoformat(),
                "method": Payment.Method.BANK,
                "reference": "ПП-CORE-001",
                "notes": "Частичная оплата",
            },
        )
        self.assertRedirects(payment_response, transportation.get_absolute_url())
        transportation.refresh_from_db()
        payment = Payment.objects.get(reference="ПП-CORE-001")
        self.assertEqual(payment.transportation_id, transportation.pk)
        self.assertIsNone(payment.shipment_id)
        self.assertEqual(transportation.receivable_balance, Decimal("45000.00"))
        self.assertEqual(
            transportation.settlement_movements.get(payment=payment).amount,
            Decimal("-15000.00"),
        )

        edit_response = self.client.post(
            reverse(
                "transportation-payment-update",
                args=[transportation.pk, payment.pk],
            ),
            {
                "direction": Payment.Direction.INCOME,
                "amount": "20000.00",
                "payment_date": document_date.isoformat(),
                "method": Payment.Method.BANK,
                "reference": "ПП-CORE-001-EDIT",
                "notes": "Изменённая частичная оплата",
            },
        )
        self.assertRedirects(edit_response, transportation.get_absolute_url())
        payment.refresh_from_db()
        self.assertEqual(payment.amount, Decimal("20000.00"))
        self.assertEqual(transportation.receivable_balance, Decimal("40000.00"))
        self.assertEqual(
            transportation.settlement_movements.get(payment=payment).amount,
            Decimal("-20000.00"),
        )

        delete_page = self.client.get(
            reverse(
                "transportation-payment-delete",
                args=[transportation.pk, payment.pk],
            )
        )
        self.assertContains(delete_page, "Удалить платёж?")
        delete_response = self.client.post(
            reverse(
                "transportation-payment-delete",
                args=[transportation.pk, payment.pk],
            )
        )
        self.assertRedirects(delete_response, transportation.get_absolute_url())
        self.assertFalse(Payment.objects.filter(pk=payment.pk).exists())
        self.assertEqual(transportation.receivable_balance, Decimal("60000.00"))
        self.assertEqual(
            transportation.status_events.filter(source="payment").count(), 3
        )

        accounting_page = self.client.get(
            reverse("transportation-accounting"), {"q": "Пластиковая"}
        )
        self.assertEqual(accounting_page.status_code, 200)
        self.assertEqual(
            accounting_page.context["accounting_totals"]["net_profit"],
            transportation.net_profit,
        )
        self.assertContains(accounting_page, "Чистая прибыль")

    def test_erp_transportation_form_and_unposting(self):
        self.client.force_login(self.user)
        create_page = self.client.get(reverse("transportation-create"))
        self.assertEqual(create_page.status_code, 200)
        self.assertContains(create_page, "Записать и провести")
        self.assertContains(create_page, "Продажа клиенту")
        self.assertContains(create_page, "Закупка у исполнителя")
        self.assertContains(create_page, "НДС к уплате")
        self.assertContains(create_page, "Прибыль рейса без НДС")
        self.assertContains(create_page, "Чистая прибыль")

        transportation = self.shipment.transportation
        transportation.customer_vat_rate = VATRate.objects.get(code="22")
        transportation.executor_vat_rate = VATRate.objects.get(code="without_vat")
        transportation.save(
            update_fields=["customer_vat_rate", "executor_vat_rate", "updated_at"]
        )
        assignment = transportation.active_vehicle_assignment()
        assignment.driver = self.driver
        assignment.vehicle = self.vehicle
        assignment.trailer_registration_number = "В456ВВ198"
        assignment.save()
        executor_contract = Contract.objects.create(
            kind=Contract.Kind.CARRIER_TRANSPORT,
            expeditor=self.company_profile,
            carrier=self.carrier,
            contract_date=date.today(),
        )
        execution_link = transportation.active_execution_link()
        execution_link.contract = executor_contract
        execution_link.save(update_fields=["contract", "updated_at"])
        posted = self.client.post(
            reverse("transportation-post", args=[transportation.pk])
        )
        self.assertRedirects(posted, transportation.get_absolute_url())
        transportation.refresh_from_db()
        self.assertEqual(
            transportation.posting_status, Transportation.PostingStatus.POSTED
        )
        unposted = self.client.post(
            reverse("transportation-unpost", args=[transportation.pk])
        )
        self.assertRedirects(
            unposted, reverse("transportation-update", args=[transportation.pk])
        )
        transportation.refresh_from_db()
        self.assertEqual(
            transportation.posting_status, Transportation.PostingStatus.DRAFT
        )
        self.assertFalse(transportation.charges.exists())
        self.assertFalse(transportation.settlement_movements.exists())

    def test_transportation_unpost_shows_message_when_accrual_is_in_reconciliation_act(self):
        self.client.force_login(self.user)
        transportation = self.shipment.transportation
        transportation.customer_vat_rate = VATRate.objects.get(code="22")
        transportation.executor_vat_rate = VATRate.objects.get(code="without_vat")
        transportation.save(
            update_fields=["customer_vat_rate", "executor_vat_rate", "updated_at"]
        )
        assignment = transportation.active_vehicle_assignment()
        assignment.driver = self.driver
        assignment.vehicle = self.vehicle
        assignment.trailer_registration_number = "В456ВВ198"
        assignment.save()
        executor_contract = Contract.objects.create(
            kind=Contract.Kind.CARRIER_TRANSPORT,
            expeditor=self.company_profile,
            carrier=self.carrier,
            contract_date=date.today(),
        )
        execution_link = transportation.active_execution_link()
        execution_link.contract = executor_contract
        execution_link.save(update_fields=["contract", "updated_at"])
        self.client.post(reverse("transportation-post", args=[transportation.pk]))
        transportation.refresh_from_db()
        accrual = transportation.settlement_movements.filter(
            kind=SettlementMovement.Kind.ACCRUAL,
            side=SettlementMovement.Side.RECEIVABLE,
        ).get()
        act = ReconciliationAct.objects.create(
            owner_company=transportation.owner_company,
            counterparty=self.customer.organization,
            period_from=date.today(),
            period_to=date.today(),
            currency="RUB",
        )
        ReconciliationActLine.objects.create(
            act=act,
            movement=accrual,
            transportation=transportation,
            movement_date=accrual.movement_date,
            description="Начисление по рейсу",
            debit=accrual.amount,
            credit=Decimal("0.00"),
            balance=accrual.amount,
        )

        response = self.client.post(
            reverse("transportation-unpost", args=[transportation.pk]),
            follow=True,
        )

        transportation.refresh_from_db()
        self.assertEqual(
            transportation.posting_status,
            Transportation.PostingStatus.POSTED,
        )
        self.assertContains(response, "Нельзя отменить проведение рейса")
        self.assertContains(response, act.number)

        void_response = self.client.post(reverse("reconciliation-act-void", args=[act.pk]))
        self.assertRedirects(void_response, reverse("reconciliation-act-list"))
        act.refresh_from_db()
        self.assertEqual(act.status, ReconciliationAct.Status.VOIDED)
        self.assertFalse(act.lines.exists())
        self.assertTrue(
            PlannerTask.objects.filter(
                title=f"Пересоздать акт сверки {act.number}",
                status=PlannerTask.Status.TODO,
            ).exists()
        )

        unpost_response = self.client.post(
            reverse("transportation-unpost", args=[transportation.pk])
        )
        self.assertRedirects(
            unpost_response, reverse("transportation-update", args=[transportation.pk])
        )
        transportation.refresh_from_db()
        self.assertEqual(
            transportation.posting_status,
            Transportation.PostingStatus.DRAFT,
        )

    def test_bank_statement_mass_post_creates_payments_for_selected_trips(self):
        self.client.force_login(self.user)
        transportation = self.shipment.transportation
        owner = self.company_profile.organization
        transportation.posting_status = Transportation.PostingStatus.POSTED
        transportation.save(update_fields=["posting_status", "updated_at"])
        SettlementMovement.objects.create(
            transportation=transportation,
            side=SettlementMovement.Side.RECEIVABLE,
            kind=SettlementMovement.Kind.ACCRUAL,
            owner_company=owner,
            counterparty=self.customer.organization,
            amount=Decimal("100000.00"),
            currency="RUB",
            movement_date=date.today(),
        )
        SettlementMovement.objects.create(
            transportation=transportation,
            side=SettlementMovement.Side.PAYABLE,
            kind=SettlementMovement.Kind.ACCRUAL,
            owner_company=owner,
            counterparty=self.carrier.organization,
            amount=Decimal("75000.00"),
            currency="RUB",
            movement_date=date.today(),
        )

        response = self.client.post(
            reverse("bank-statement-create"),
            {
                "direction": BankStatement.Direction.INCOME,
                "statement_date": date.today().isoformat(),
                "owner_company": owner.pk,
                "bank_account": "",
                "currency": "RUB",
                "reference": "ВЫП-001",
                "notes": "Массовая обработка",
                "transportation_ids": str(transportation.pk),
                f"reference_{transportation.pk}": "ПП-ВЫП-001",
                f"amount_{transportation.pk}": "40000.00",
                "action": "post",
            },
        )
        statement = BankStatement.objects.get(reference="ВЫП-001")
        self.assertRedirects(response, statement.get_absolute_url())
        self.assertEqual(statement.status, BankStatement.Status.POSTED)
        line = statement.lines.get()
        self.assertEqual(line.payment_reference, "ПП-ВЫП-001")
        self.assertEqual(line.payment.direction, Payment.Direction.INCOME)
        self.assertEqual(line.payment.reference, "ПП-ВЫП-001")
        self.assertEqual(line.payment.amount, Decimal("40000.00"))
        transportation_page = self.client.get(transportation.get_absolute_url())
        self.assertContains(transportation_page, "Оплачено")
        self.assertContains(transportation_page, "ПП-ВЫП-001")
        transportation.refresh_from_db()
        self.assertEqual(transportation.receivable_balance, Decimal("60000.00"))

        unpost_response = self.client.post(
            reverse("bank-statement-unpost", args=[statement.pk])
        )
        self.assertRedirects(
            unpost_response, reverse("bank-statement-update", args=[statement.pk])
        )
        statement.refresh_from_db()
        self.assertEqual(statement.status, BankStatement.Status.DRAFT)
        self.assertFalse(Payment.objects.filter(reference="ПП-ВЫП-001").exists())
        transportation.refresh_from_db()
        self.assertEqual(transportation.receivable_balance, Decimal("100000.00"))
        self.client.post(reverse("bank-statement-post", args=[statement.pk]))
        statement.refresh_from_db()
        line.refresh_from_db()
        self.assertEqual(statement.status, BankStatement.Status.POSTED)
        self.assertTrue(line.payment_id)

        line_update_response = self.client.post(
            reverse(
                "bank-statement-line-update",
                args=[statement.pk, line.pk],
            ),
            {"amount": "45000.00", "payment_reference": "ПП-ВЫП-001-ИЗМ"},
        )
        self.assertRedirects(line_update_response, statement.get_absolute_url())
        line.refresh_from_db()
        line.payment.refresh_from_db()
        self.assertEqual(line.amount, Decimal("45000.00"))
        self.assertEqual(line.payment.reference, "ПП-ВЫП-001-ИЗМ")
        transportation.refresh_from_db()
        self.assertEqual(transportation.receivable_balance, Decimal("55000.00"))

        line_delete_response = self.client.post(
            reverse("bank-statement-line-delete", args=[statement.pk, line.pk])
        )
        self.assertRedirects(line_delete_response, statement.get_absolute_url())
        self.assertFalse(BankStatementLine.objects.filter(pk=line.pk).exists())
        transportation.refresh_from_db()
        self.assertEqual(transportation.receivable_balance, Decimal("100000.00"))

        expense_response = self.client.post(
            reverse("bank-statement-create"),
            {
                "direction": BankStatement.Direction.EXPENSE,
                "statement_date": date.today().isoformat(),
                "owner_company": owner.pk,
                "bank_account": "",
                "currency": "RUB",
                "reference": "ВЫП-002",
                "notes": "Оплата исполнителю",
                "transportation_ids": str(transportation.pk),
                f"amount_{transportation.pk}": "25000.00",
                "action": "post",
            },
        )
        expense = BankStatement.objects.get(reference="ВЫП-002")
        self.assertRedirects(expense_response, expense.get_absolute_url())
        self.assertEqual(expense.status, BankStatement.Status.POSTED)
        self.assertEqual(expense.lines.get().payment.direction, Payment.Direction.EXPENSE)
        transportation.refresh_from_db()
        self.assertEqual(transportation.payable_balance, Decimal("50000.00"))

        delete_response = self.client.post(
            reverse("bank-statement-delete", args=[expense.pk])
        )
        self.assertRedirects(delete_response, reverse("bank-statement-list"))
        self.assertFalse(BankStatement.objects.filter(pk=expense.pk).exists())
        transportation.refresh_from_db()
        self.assertEqual(transportation.payable_balance, Decimal("75000.00"))

    def test_bank_statement_unpost_shows_message_when_payment_is_in_reconciliation_act(self):
        self.client.force_login(self.user)
        transportation = self.shipment.transportation
        owner = self.company_profile.organization
        transportation.posting_status = Transportation.PostingStatus.POSTED
        transportation.save(update_fields=["posting_status", "updated_at"])
        accrual = SettlementMovement.objects.create(
            transportation=transportation,
            side=SettlementMovement.Side.RECEIVABLE,
            kind=SettlementMovement.Kind.ACCRUAL,
            owner_company=owner,
            counterparty=self.customer.organization,
            amount=Decimal("100000.00"),
            currency="RUB",
            movement_date=date.today(),
        )
        response = self.client.post(
            reverse("bank-statement-create"),
            {
                "direction": BankStatement.Direction.INCOME,
                "statement_date": date.today().isoformat(),
                "owner_company": owner.pk,
                "bank_account": "",
                "currency": "RUB",
                "reference": "ВЫП-АС-001",
                "notes": "",
                "transportation_ids": str(transportation.pk),
                f"reference_{transportation.pk}": "ПП-АС-001",
                f"amount_{transportation.pk}": "40000.00",
                "action": "post",
            },
        )
        statement = BankStatement.objects.get(reference="ВЫП-АС-001")
        self.assertRedirects(response, statement.get_absolute_url())
        payment_movement = statement.lines.get().payment.settlement_movement
        act = ReconciliationAct.objects.create(
            owner_company=owner,
            counterparty=self.customer.organization,
            period_from=date.today(),
            period_to=date.today(),
            currency="RUB",
        )
        ReconciliationActLine.objects.create(
            act=act,
            movement=payment_movement,
            transportation=transportation,
            movement_date=payment_movement.movement_date,
            description="Оплата по выписке",
            debit=Decimal("0.00"),
            credit=Decimal("40000.00"),
            balance=Decimal("60000.00"),
        )
        ReconciliationActLine.objects.create(
            act=act,
            movement=accrual,
            transportation=transportation,
            movement_date=accrual.movement_date,
            description="Начисление",
            debit=Decimal("100000.00"),
            credit=Decimal("0.00"),
            balance=Decimal("100000.00"),
        )

        unpost_response = self.client.post(
            reverse("bank-statement-unpost", args=[statement.pk]),
            follow=True,
        )

        statement.refresh_from_db()
        self.assertEqual(statement.status, BankStatement.Status.POSTED)
        self.assertContains(unpost_response, "Нельзя отменить платёж")
        self.assertContains(unpost_response, act.number)
        self.assertTrue(Payment.objects.filter(reference="ПП-АС-001").exists())

        void_response = self.client.post(reverse("reconciliation-act-void", args=[act.pk]))
        self.assertRedirects(void_response, reverse("reconciliation-act-list"))
        act.refresh_from_db()
        self.assertEqual(act.status, ReconciliationAct.Status.VOIDED)
        self.assertFalse(act.lines.exists())
        self.assertTrue(
            PlannerTask.objects.filter(
                title=f"Пересоздать акт сверки {act.number}",
                status=PlannerTask.Status.TODO,
            ).exists()
        )

        unpost_after_void = self.client.post(
            reverse("bank-statement-unpost", args=[statement.pk])
        )
        self.assertRedirects(
            unpost_after_void, reverse("bank-statement-update", args=[statement.pk])
        )
        statement.refresh_from_db()
        self.assertEqual(statement.status, BankStatement.Status.DRAFT)
        self.assertFalse(Payment.objects.filter(reference="ПП-АС-001").exists())

    def test_pages_use_uikit_components(self):
        self.client.force_login(self.user)
        dashboard = self.client.get(reverse("dashboard"))
        self.assertContains(dashboard, "uikit@3.25.21")
        self.assertContains(dashboard, "v=20260908-directory-unified")
        self.assertContains(dashboard, "uikit-theme")
        self.assertContains(dashboard, "uk-card uk-card-default")
        self.assertContains(dashboard, "Dashboard")
        self.assertContains(dashboard, "dashboard-today")
        shipments = self.client.get(reverse("shipment-list"))
        self.assertContains(shipments, "uk-table uk-table-small")
        form = self.client.get(reverse("shipment-update", args=[self.shipment.pk]))
        self.assertContains(form, 'class="form-control uk-input"')
        self.assertContains(form, 'class="form-control uk-select"')
        drivers = self.client.get(reverse("driver-list"))
        self.assertContains(drivers, "driver-directory-table")
        self.assertNotContains(drivers, "bootstrap@5")
        driver_form = self.client.get(reverse("driver-create"))
        self.assertContains(driver_form, "driver-form-workspace")
        self.assertNotContains(driver_form, "bootstrap@5")

    def test_driver_and_vehicle_profiles_contain_operational_information(self):
        self.client.force_login(self.user)
        driver_response = self.client.get(reverse("driver-detail", args=[self.driver.pk]))
        self.assertContains(driver_response, self.driver.license_number)
        self.assertContains(driver_response, self.driver.license_categories)
        vehicle_response = self.client.get(
            reverse("vehicle-detail", args=[self.vehicle.pk])
        )
        self.assertContains(vehicle_response, self.vehicle.registration_number)
        self.assertContains(vehicle_response, self.vehicle.make)

    def test_driver_directory_and_card_follow_counterparty_workspace(self):
        self.client.force_login(self.user)
        listing = self.client.get(
            reverse("driver-list"),
            {"q": self.driver.last_name, "active": "1", "carrier": self.carrier.pk},
        )
        self.assertEqual(listing.context["page_obj"].paginator.count, 1)
        self.assertContains(listing, "driver-directory-table")
        self.assertContains(listing, "Документы требуют внимания")
        detail = self.client.get(reverse("driver-detail", args=[self.driver.pk]))
        self.assertContains(detail, "onec-titlebar")
        self.assertContains(detail, "driver-summary-card")
        self.assertContains(detail, "Компании водителя")
        self.assertContains(detail, "История карточки")

    def test_vehicle_directory_and_card_follow_onec_workspace(self):
        self.client.force_login(self.user)
        listing = self.client.get(
            reverse("vehicle-list"),
            {
                "q": self.vehicle.registration_number.lower(),
                "active": "1",
                "kind": Vehicle.Kind.TRACTOR,
                "carrier": self.carrier.pk,
            },
        )
        self.assertEqual(listing.context["page_obj"].paginator.count, 1)
        self.assertContains(listing, "vehicle-directory-table")
        self.assertContains(listing, "Документы требуют внимания")
        detail = self.client.get(reverse("vehicle-detail", args=[self.vehicle.pk]))
        self.assertContains(detail, "onec-titlebar")
        self.assertContains(detail, "vehicle-summary-card")
        self.assertContains(detail, "Сцепки с прицепами")
        self.assertContains(detail, "История карточки")

    def test_vehicle_combination_is_separate_and_supports_gazelle_without_trailer(self):
        self.client.force_login(self.user)
        trailer = Vehicle.objects.create(
            carrier=self.carrier,
            kind=Vehicle.Kind.SEMITRAILER,
            registration_number="П001ЕСТ",
            make="Schmitz",
            model="S.KO",
        )
        form = VehicleCombinationForm(
            data={
                "tractor": self.vehicle.pk,
                "trailer": trailer.pk,
                "valid_from": date.today().isoformat(),
                "valid_until": "",
                "is_active": "on",
                "notes": "Основная сцепка",
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        combination = form.save()
        self.assertEqual(combination.tractor, self.vehicle)
        self.assertEqual(combination.trailer, trailer)

        duplicate = VehicleCombinationForm(
            data={
                "tractor": self.vehicle.pk,
                "trailer": trailer.pk,
                "valid_from": "",
                "valid_until": "",
                "is_active": "on",
                "notes": "Дубликат",
            }
        )
        self.assertFalse(duplicate.is_valid())
        self.assertIn("активная сцепка", duplicate.errors.as_text())

        api_response = self.client.get(
            reverse("organization-resources"),
            {"organization": self.carrier.organization.pk},
        )
        self.assertEqual(api_response.status_code, 200)
        self.assertEqual(api_response.json()["combinations"][0]["id"], combination.pk)
        self.assertEqual(
            api_response.json()["combinations"][0]["trailer_id"], trailer.pk
        )

        vehicle_response = self.client.get(
            reverse("vehicle-detail", args=[self.vehicle.pk])
        )
        self.assertContains(vehicle_response, "Сцепки с прицепами")
        self.assertContains(vehicle_response, trailer.registration_number)

        chain_form = TransportationChainForm(
            data={
                "executor": self.carrier.organization.pk,
                "executor_role": TransportationLink.ContractorRole.CARRIER,
                "contract": "",
                "instruction_number": "",
                "instruction_status": "",
                "actual_carrier": self.carrier.organization.pk,
                "driver": self.driver.pk,
                "vehicle": self.vehicle.pk,
                "combination": combination.pk,
                "trailer": trailer.pk,
                "trailer_registration_number": "",
            },
            transportation=self.shipment.transportation,
        )
        self.assertTrue(chain_form.is_valid(), chain_form.errors)
        chain_form.save(self.user)
        assignment = self.shipment.transportation.active_vehicle_assignment()
        self.assertEqual(assignment.combination, combination)
        self.assertEqual(assignment.trailer, trailer)

        gazelle = Vehicle.objects.create(
            carrier=self.carrier,
            kind=Vehicle.Kind.GAZELLE,
            registration_number="Г001ЕСТ",
            make="ГАЗ",
            model="Газель Next",
        )
        self.assertFalse(gazelle.requires_trailer)
        self.assertFalse(gazelle.can_tow_trailer)
        assignment.vehicle = gazelle
        assignment.trailer = None
        assignment.combination = None
        assignment.save()
        self.assertFalse(
            any("полуприцеп" in issue for issue in self.shipment.transportation.chain_issues())
        )
        assignment.trailer = trailer
        with self.assertRaises(ValidationError) as error:
            assignment.full_clean()
        self.assertIn("trailer", error.exception.message_dict)

    def test_shipment_resources_must_belong_to_selected_carrier(self):
        other_carrier = Carrier.objects.create(name="Другой перевозчик")
        other_driver = Driver.objects.create(
            carrier=other_carrier,
            last_name="Петров",
            first_name="Пётр",
            phone="+7 900 000-00-02",
            license_number="TEST-DRIVER-2",
            license_categories="C, CE",
            license_expiry_date=date.today() + timedelta(days=365),
        )
        other_vehicle = Vehicle.objects.create(
            carrier=other_carrier,
            kind=Vehicle.Kind.TRUCK,
            registration_number="Т002ЕСТ",
            make="MAN",
        )
        self.shipment.carrier = self.carrier
        self.shipment.driver = other_driver
        self.shipment.vehicle = other_vehicle
        with self.assertRaises(ValidationError) as error:
            self.shipment.full_clean()
        self.assertIn("driver", error.exception.message_dict)
        self.assertIn("vehicle", error.exception.message_dict)

    def test_shipment_accepts_resources_of_selected_carrier(self):
        self.shipment.driver = self.driver
        self.shipment.vehicle = self.vehicle
        self.shipment.full_clean()
        self.shipment.save()
        self.shipment.refresh_from_db()
        self.assertEqual(self.shipment.driver, self.driver)
        self.assertEqual(self.shipment.vehicle, self.vehicle)

    def test_carrier_resources_api_filters_drivers_and_vehicles(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("carrier-resources"), {"carrier": self.carrier.pk}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["drivers"][0]["id"], self.driver.pk)
        self.assertEqual(response.json()["vehicles"][0]["id"], self.vehicle.pk)

    def test_dadata_lookup_requires_login(self):
        response = self.client.post(
            reverse("dadata-party-by-inn"), {"inn": "7707083893"}
        )
        self.assertEqual(response.status_code, 302)

    def test_dadata_address_suggestions_require_login(self):
        response = self.client.get(
            reverse("dadata-address-suggestions"), {"q": "Тверская 1"}
        )
        self.assertEqual(response.status_code, 302)

    def test_dadata_bank_lookup_requires_login(self):
        response = self.client.get(reverse("dadata-bank-by-bik"), {"bik": "044525225"})
        self.assertEqual(response.status_code, 302)

    def test_dadata_address_suggestions_ignore_short_queries(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("dadata-address-suggestions"), {"q": "Тв"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"suggestions": []})

    @override_settings(
        DADATA_API_TOKEN="test-server-token",
        DADATA_ADDRESS_URL="https://suggestions.test/suggest/address",
    )
    @patch("crm.dadata.urlopen")
    def test_dadata_address_suggestions_are_proxied_and_normalized(
        self, mocked_urlopen
    ):
        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def read(self):
                return json.dumps(
                    {
                        "suggestions": [
                            {
                                "value": "г Москва, ул Тверская, д 1",
                                "unrestricted_value": (
                                    "125009, г Москва, ул Тверская, д 1"
                                ),
                                "data": {
                                    "postal_code": "125009",
                                    "region_with_type": "г Москва",
                                    "city_with_type": "г Москва",
                                    "fias_id": "test-fias-id",
                                },
                            }
                        ]
                    },
                    ensure_ascii=False,
                ).encode("utf-8")

        mocked_urlopen.return_value = FakeResponse()
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("dadata-address-suggestions"),
            {"q": "Тверская 1", "city": "Москва"},
        )

        self.assertEqual(response.status_code, 200)
        suggestion = response.json()["suggestions"][0]
        self.assertEqual(suggestion["value"], "г Москва, ул Тверская, д 1")
        self.assertEqual(suggestion["postal_code"], "125009")
        self.assertEqual(suggestion["fias_id"], "test-fias-id")
        request = mocked_urlopen.call_args.args[0]
        self.assertEqual(request.full_url, "https://suggestions.test/suggest/address")
        self.assertEqual(request.get_header("Authorization"), "Token test-server-token")
        self.assertEqual(
            json.loads(request.data),
            {"query": "Москва, Тверская 1", "count": 10},
        )

    def test_dadata_lookup_validates_inn(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("dadata-party-by-inn"), {"inn": "123"}
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("10 цифр", response.json()["error"])

    def test_dadata_bank_lookup_validates_bik(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("dadata-bank-by-bik"), {"bik": "044"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("9 цифр", response.json()["error"])

    @override_settings(
        DADATA_API_TOKEN="test-server-token",
        DADATA_BANK_URL="https://suggestions.test/findById/bank",
    )
    @patch("crm.dadata.urlopen")
    def test_dadata_bank_lookup_returns_normalized_bank_details(self, mocked_urlopen):
        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def read(self):
                return json.dumps(
                    {
                        "suggestions": [
                            {
                                "value": "ПАО СБЕРБАНК",
                                "data": {
                                    "bic": "044525225",
                                    "correspondent_account": "30101810400000000225",
                                    "name": {
                                        "payment": "ПАО СБЕРБАНК",
                                        "short": "СБЕРБАНК",
                                    },
                                    "address": {"value": "г Москва"},
                                },
                            }
                        ]
                    },
                    ensure_ascii=False,
                ).encode("utf-8")

        mocked_urlopen.return_value = FakeResponse()
        self.client.force_login(self.user)
        response = self.client.get(reverse("dadata-bank-by-bik"), {"bik": "044525225"})

        self.assertEqual(response.status_code, 200)
        bank = response.json()["bank"]
        self.assertEqual(bank["bank_name"], "ПАО СБЕРБАНК")
        self.assertEqual(bank["bik"], "044525225")
        self.assertEqual(bank["correspondent_account"], "30101810400000000225")
        request = mocked_urlopen.call_args.args[0]
        self.assertEqual(request.full_url, "https://suggestions.test/findById/bank")
        self.assertEqual(json.loads(request.data), {"query": "044525225", "count": 1})

    @override_settings(DADATA_API_TOKEN="test-server-token")
    @patch("crm.dadata.urlopen")
    def test_dadata_lookup_returns_normalized_company_details(self, mocked_urlopen):
        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def read(self):
                return json.dumps(
                    {
                        "suggestions": [
                            {
                                "value": 'ООО "МОТОРИКА"',
                                "data": {
                                    "inn": "7707083893",
                                    "kpp": "772301001",
                                    "ogrn": "1027700132195",
                                    "type": "LEGAL",
                                    "name": {
                                        "full_with_opf": 'Общество с ограниченной ответственностью "МОТОРИКА"',
                                        "short_with_opf": 'ООО "МОТОРИКА"',
                                    },
                                    "address": {
                                        "value": "г Москва, ул Тестовая, д 1",
                                        "unrestricted_value": "123456, г Москва, ул Тестовая, д 1",
                                    },
                                    "management": {
                                        "name": "Иванов Иван Иванович",
                                        "post": "ГЕНЕРАЛЬНЫЙ ДИРЕКТОР",
                                    },
                                    "phones": [{"value": "+7 495 000-00-00"}],
                                    "emails": [{"value": "info@example.com"}],
                                    "state": {"status": "ACTIVE"},
                                    "invalid": None,
                                },
                            }
                        ]
                    },
                    ensure_ascii=False,
                ).encode("utf-8")

        mocked_urlopen.return_value = FakeResponse()
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("dadata-party-by-inn"), {"inn": "7707083893"}
        )
        self.assertEqual(response.status_code, 200)
        party = response.json()["party"]
        self.assertEqual(party["short_name"], 'ООО "МОТОРИКА"')
        self.assertEqual(party["kpp"], "772301001")
        self.assertEqual(party["director_name"], "Иванов Иван Иванович")
        self.assertEqual(party["status_label"], "Действующая")
        request = mocked_urlopen.call_args.args[0]
        self.assertEqual(request.get_header("Authorization"), "Token test-server-token")
        self.assertEqual(json.loads(request.data)["branch_type"], "MAIN")

    @override_settings(DADATA_API_TOKEN="")
    def test_dadata_lookup_explains_missing_server_token(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("dadata-party-by-inn"), {"inn": "7812345671"}
        )
        self.assertEqual(response.status_code, 503)
        self.assertIn("DADATA_API_TOKEN", response.json()["error"])

    def test_company_forms_offer_dadata_autofill(self):
        self.client.force_login(self.user)
        for url in (
            reverse("expeditor-create"),
            reverse("customer-create"),
            reverse("carrier-create"),
        ):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertContains(response, "Заполнение по ИНН через DaData")
                self.assertContains(response, reverse("dadata-party-by-inn"))

    def test_route_address_fields_offer_dadata_suggestions(self):
        self.client.force_login(self.user)
        for url in (reverse("transportation-create"), reverse("shipment-create")):
            with self.subTest(url=url):
                response = self.client.get(url)
                expected_address_fields = 3 if url == reverse("transportation-create") else 2
                expected_city_fields = 3 if url == reverse("transportation-create") else 2
                self.assertContains(
                    response,
                    'data-dadata-address=""',
                    count=expected_address_fields,
                )
                self.assertContains(
                    response,
                    'data-dadata-city=""',
                    count=expected_city_fields,
                )
                self.assertContains(response, reverse("dadata-address-suggestions"))
                self.assertContains(response, "js/address-suggestions.js")
                self.assertContains(response, "js/city-suggestions.js")

    @override_settings(
        DADATA_API_TOKEN="test-server-token",
        DADATA_FIO_URL="https://suggestions.test/suggest/fio",
    )
    @patch("crm.dadata.urlopen")
    def test_dadata_fio_suggestions_are_proxied_and_normalized(
        self, mocked_urlopen
    ):
        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def read(self):
                return json.dumps(
                    {
                        "suggestions": [
                            {
                                "value": "Иванов",
                                "data": {
                                    "surname": "Иванов",
                                    "name": None,
                                    "patronymic": None,
                                    "gender": "MALE",
                                },
                            }
                        ]
                    },
                    ensure_ascii=False,
                ).encode("utf-8")

        mocked_urlopen.return_value = FakeResponse()
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("dadata-fio-suggestions"),
            {"q": "Иван", "part": "SURNAME"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["suggestions"][0]["surname"], "Иванов")
        request = mocked_urlopen.call_args.args[0]
        self.assertEqual(request.full_url, "https://suggestions.test/suggest/fio")
        self.assertEqual(
            json.loads(request.data),
            {"query": "Иван", "count": 10, "parts": ["SURNAME"]},
        )

    @override_settings(
        DADATA_API_TOKEN="test-server-token",
        DADATA_FMS_UNIT_URL="https://suggestions.test/suggest/fms-unit",
    )
    @patch("crm.dadata.urlopen")
    def test_dadata_fms_unit_suggestions_are_proxied_and_normalized(
        self, mocked_urlopen
    ):
        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def read(self):
                return json.dumps(
                    {
                        "suggestions": [
                            {
                                "value": "ОВД района Тверской г. Москвы",
                                "data": {
                                    "code": "770-001",
                                    "name": "ОВД района Тверской г. Москвы",
                                    "region_code": "77",
                                    "type": "2",
                                },
                            }
                        ]
                    },
                    ensure_ascii=False,
                ).encode("utf-8")

        mocked_urlopen.return_value = FakeResponse()
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("dadata-fms-unit-suggestions"), {"q": "770-001"}
        )

        self.assertEqual(response.status_code, 200)
        suggestion = response.json()["suggestions"][0]
        self.assertEqual(suggestion["code"], "770-001")
        self.assertEqual(suggestion["value"], "ОВД района Тверской г. Москвы")
        request = mocked_urlopen.call_args.args[0]
        self.assertEqual(
            request.full_url, "https://suggestions.test/suggest/fms-unit"
        )
        self.assertEqual(json.loads(request.data), {"query": "770-001", "count": 10})

    def test_driver_forms_hide_removed_fields_and_offer_dadata_suggestions(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("driver-create"))
        form = response.context["form"]
        for field_name in ("email", "address", "medical_certificate_expiry"):
            self.assertNotIn(field_name, form.fields)
        for field_name in ("last_name", "first_name", "middle_name"):
            self.assertEqual(
                form.fields[field_name].widget.attrs["data-dadata-driver"], "fio"
            )
        passport_form = response.context["passport_formset"].forms[0]
        self.assertEqual(
            passport_form.fields["issued_by"].widget.attrs["data-dadata-driver"],
            "fms",
        )
        self.assertNotIn("additional_carriers", form.fields)
        self.assertEqual(form.fields["carrier"].widget.input_type, "hidden")
        self.assertEqual(form.fields["license_number"].widget.input_type, "hidden")
        self.assertIn("tax_id", form.fields)
        self.assertIn("employment_formset", response.context)
        self.assertIn("license_formset", response.context)
        self.assertContains(response, "Добавить перевозчика")
        self.assertContains(response, "Добавить ещё паспорт")
        self.assertContains(response, "Добавить ещё удостоверение")
        self.assertContains(response, "js/driver-suggestions.js")

    def test_driver_create_opens_in_modal_and_saves_inn_and_passport(self):
        self.client.force_login(self.user)
        additional_carrier_one = Carrier.objects.create(name="Доп. перевозчик 1")
        additional_carrier_two = Carrier.objects.create(name="Доп. перевозчик 2")
        listing = self.client.get(reverse("driver-list"))
        self.assertContains(listing, "data-driver-modal")
        self.assertContains(listing, "js/driver-modal.js")

        response = self.client.post(
            reverse("driver-create"),
            {
                "last_name": "Сидоров",
                "first_name": "Семён",
                "phone": "+7 900 333-44-55",
                "tax_id": "7812 3456 7890",
                "is_active": "on",
                "employments-TOTAL_FORMS": "3",
                "employments-INITIAL_FORMS": "0",
                "employments-MIN_NUM_FORMS": "0",
                "employments-MAX_NUM_FORMS": "1000",
                "employments-0-carrier": str(self.carrier.pk),
                "employments-0-is_primary": "on",
                "employments-1-carrier": str(additional_carrier_one.pk),
                "employments-2-carrier": str(additional_carrier_two.pk),
                "licenses-TOTAL_FORMS": "1",
                "licenses-INITIAL_FORMS": "0",
                "licenses-MIN_NUM_FORMS": "0",
                "licenses-MAX_NUM_FORMS": "1000",
                "licenses-0-number": "MODAL-DRIVER",
                "licenses-0-categories": "C, CE",
                "licenses-0-expiry_date": "31.08.2030",
                "licenses-0-is_current": "on",
                "passports-TOTAL_FORMS": "1",
                "passports-INITIAL_FORMS": "0",
                "passports-MIN_NUM_FORMS": "0",
                "passports-MAX_NUM_FORMS": "1000",
                "passports-0-series": "4010",
                "passports-0-number": "987654",
                "passports-0-issued_by": "ОВД района Тверской",
                "passports-0-issue_date": "01.09.2025",
                "passports-0-is_current": "on",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
            HTTP_ACCEPT="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        driver = Driver.objects.get(license_number="MODAL-DRIVER")
        self.assertEqual(driver.tax_id, "781234567890")
        self.assertTrue(
            driver.passports.filter(
                series="4010", number="987654", is_current=True
            ).exists()
        )
        self.assertTrue(
            driver.employments.filter(
                carrier=self.carrier, is_primary=True, is_active=True
            ).exists()
        )
        self.assertEqual(
            set(
                driver.employments.filter(
                    is_active=True, is_primary=False
                ).values_list("carrier_id", flat=True)
            ),
            {additional_carrier_one.pk, additional_carrier_two.pk},
        )
        self.assertTrue(
            driver.licenses.filter(
                number="MODAL-DRIVER", categories="C, CE", is_current=True
            ).exists()
        )
        edit_response = self.client.get(reverse("driver-update", args=[driver.pk]))
        self.assertEqual(
            {
                row.instance.carrier_id
                for row in edit_response.context["employment_formset"].forms
                if row.instance.pk
            },
            {self.carrier.pk, additional_carrier_one.pk, additional_carrier_two.pk},
        )
        self.assertContains(edit_response, "data-employment-add")
        self.assertContains(edit_response, "data-license-add")

    def test_driver_can_be_created_without_license_and_registry_prompts_to_add_it(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("driver-create"),
            {
                "last_name": "Безправов",
                "first_name": "Павел",
                "phone": "+7 900 444-55-66",
                "tax_id": "",
                "is_active": "on",
                "employments-TOTAL_FORMS": "1",
                "employments-INITIAL_FORMS": "0",
                "employments-MIN_NUM_FORMS": "0",
                "employments-MAX_NUM_FORMS": "1000",
                "employments-0-carrier": str(self.carrier.pk),
                "employments-0-is_primary": "on",
                "licenses-TOTAL_FORMS": "1",
                "licenses-INITIAL_FORMS": "0",
                "licenses-MIN_NUM_FORMS": "0",
                "licenses-MAX_NUM_FORMS": "1000",
                "licenses-0-number": "",
                "licenses-0-categories": "",
                "licenses-0-issue_date": "",
                "licenses-0-expiry_date": "",
                "passports-TOTAL_FORMS": "0",
                "passports-INITIAL_FORMS": "0",
                "passports-MIN_NUM_FORMS": "0",
                "passports-MAX_NUM_FORMS": "1000",
            },
        )

        driver = Driver.objects.get(last_name="Безправов")
        self.assertRedirects(response, driver.get_absolute_url())
        self.assertEqual(driver.license_number, "")
        self.assertIsNone(driver.license_expiry_date)
        self.assertFalse(driver.licenses.exists())

        listing = self.client.get(reverse("driver-list"), {"q": "Безправов"})
        self.assertContains(listing, "Внесите ВУ")
        detail = self.client.get(driver.get_absolute_url())
        self.assertContains(detail, "Внесите ВУ")

    def test_driver_passport_history_keeps_one_current_and_syncs_legacy_fields(self):
        first = DriverPassport.objects.create(
            driver=self.driver,
            series="4501",
            number="123456",
            issued_by="ОВД района Тверской",
            issue_date=date(2015, 1, 10),
            is_current=True,
        )
        second = DriverPassport.objects.create(
            driver=self.driver,
            series="4502",
            number="654321",
            issued_by="ОВД района Арбат",
            issue_date=date(2025, 2, 20),
            is_current=True,
        )

        first.refresh_from_db()
        self.driver.refresh_from_db()
        self.assertFalse(first.is_current)
        self.assertTrue(second.is_current)
        self.assertEqual(self.driver.passports.count(), 2)
        self.assertEqual(self.driver.passport_series, "4502")
        self.assertEqual(self.driver.passport_number, "654321")

    def test_driver_passport_series_and_number_are_globally_unique(self):
        DriverPassport.objects.create(
            driver=self.driver,
            series="45 01",
            number="123-456",
        )
        other_carrier = Carrier.objects.create(name="Второй перевозчик")
        other_driver = Driver.objects.create(
            carrier=other_carrier,
            last_name="Петров",
            first_name="Пётр",
            phone="+7 900 000-00-02",
            license_number="TEST-UNIQUE-PASSPORT",
            license_categories="C, CE",
            license_expiry_date=date.today() + timedelta(days=365),
        )

        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                DriverPassport.objects.create(
                    driver=other_driver,
                    series="4501",
                    number="123456",
                )

    def test_driver_license_history_keeps_one_current_and_syncs_legacy_fields(self):
        first = self.driver.current_license
        second = DriverLicense.objects.create(
            driver=self.driver,
            number="77 11 654321",
            categories="C, CE",
            issue_date=date(2026, 1, 10),
            expiry_date=date(2036, 1, 10),
            is_current=True,
        )

        first.refresh_from_db()
        self.driver.refresh_from_db()
        self.assertFalse(first.is_current)
        self.assertTrue(second.is_current)
        self.assertEqual(self.driver.licenses.count(), 2)
        self.assertEqual(self.driver.license_number, "77 11 654321")
        self.assertEqual(self.driver.license_expiry_date, date(2036, 1, 10))

    def test_driver_license_number_is_globally_unique_after_normalization(self):
        other_carrier = Carrier.objects.create(name="Перевозчик удостоверения")
        other_driver = Driver.objects.create(
            carrier=other_carrier,
            last_name="Петров",
            first_name="Пётр",
            phone="+7 900 000-00-03",
            license_number="UNIQUE-LICENSE-2",
            license_categories="C",
            license_expiry_date=date.today() + timedelta(days=365),
        )

        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                DriverLicense.objects.create(
                    driver=other_driver,
                    number="T E S T - DRIVER - 1",
                    categories="C",
                    expiry_date=date.today() + timedelta(days=700),
                    is_current=True,
                )

    def test_driver_update_can_switch_primary_carrier_and_current_license(self):
        self.client.force_login(self.user)
        second_carrier = Carrier.objects.create(
            name="Новый основной перевозчик", tax_id="7812345678"
        )
        employment = self.driver.employments.get(carrier=self.carrier)
        old_license = self.driver.current_license

        response = self.client.post(
            reverse("driver-update", args=[self.driver.pk]),
            {
                "last_name": self.driver.last_name,
                "first_name": self.driver.first_name,
                "phone": self.driver.phone,
                "tax_id": "",
                "is_active": "on",
                "employments-TOTAL_FORMS": "2",
                "employments-INITIAL_FORMS": "1",
                "employments-MIN_NUM_FORMS": "0",
                "employments-MAX_NUM_FORMS": "1000",
                "employments-0-id": str(employment.pk),
                "employments-0-carrier": str(self.carrier.pk),
                "employments-1-carrier": str(second_carrier.pk),
                "employments-1-is_primary": "on",
                "licenses-TOTAL_FORMS": "2",
                "licenses-INITIAL_FORMS": "1",
                "licenses-MIN_NUM_FORMS": "0",
                "licenses-MAX_NUM_FORMS": "1000",
                "licenses-0-id": str(old_license.pk),
                "licenses-0-number": old_license.number,
                "licenses-0-categories": old_license.categories,
                "licenses-0-expiry_date": old_license.expiry_date.strftime("%d.%m.%Y"),
                "licenses-1-number": "77 22 654321",
                "licenses-1-categories": "C, CE",
                "licenses-1-issue_date": "01.09.2026",
                "licenses-1-expiry_date": "01.09.2036",
                "licenses-1-is_current": "on",
                "passports-TOTAL_FORMS": "0",
                "passports-INITIAL_FORMS": "0",
                "passports-MIN_NUM_FORMS": "0",
                "passports-MAX_NUM_FORMS": "1000",
            },
        )

        self.assertRedirects(response, self.driver.get_absolute_url())
        self.driver.refresh_from_db()
        old_license.refresh_from_db()
        self.assertEqual(self.driver.carrier, second_carrier)
        self.assertFalse(old_license.is_current)
        self.assertEqual(self.driver.license_number, "77 22 654321")
        self.assertEqual(
            set(
                self.driver.employments.filter(is_active=True).values_list(
                    "carrier_id", flat=True
                )
            ),
            {self.carrier.pk, second_carrier.pk},
        )
        self.assertTrue(
            self.driver.employments.filter(
                carrier=second_carrier, is_primary=True
            ).exists()
        )

    def test_driver_can_work_for_several_carriers(self):
        second_carrier = Carrier.objects.create(name="Второй перевозчик")
        DriverEmployment.objects.create(
            driver=self.driver,
            carrier=second_carrier,
            is_active=True,
        )
        self.shipment.carrier = second_carrier
        self.shipment.driver = self.driver
        self.shipment.full_clean()

        self.client.force_login(self.user)
        response = self.client.get(
            reverse("carrier-resources"), {"carrier": second_carrier.pk}
        )
        self.assertIn(
            self.driver.pk,
            [item["id"] for item in response.json()["drivers"]],
        )

    def test_driver_pages_do_not_show_removed_personal_fields(self):
        self.driver.email = "hidden-driver@example.test"
        self.driver.address = "Скрытый адрес регистрации"
        self.driver.medical_certificate_expiry = date(2030, 1, 1)
        self.driver.save()
        self.client.force_login(self.user)

        detail = self.client.get(reverse("driver-detail", args=[self.driver.pk]))
        listing = self.client.get(reverse("driver-list"))
        for response in (detail, listing):
            self.assertNotContains(response, "hidden-driver@example.test")
            self.assertNotContains(response, "Скрытый адрес регистрации")
        self.assertNotContains(detail, "Медицинская справка")

    def test_contract_requires_counterparty_for_selected_template(self):
        contract = Contract(
            kind=Contract.Kind.CLIENT_FORWARDING,
            expeditor=self.company_profile,
            contract_date=date.today(),
        )
        with self.assertRaises(ValidationError) as error:
            contract.full_clean()
        self.assertIn("customer", error.exception.message_dict)

        contract.customer = self.customer
        contract.full_clean()

    def test_terminated_contract_requires_termination_date(self):
        contract = Contract(
            kind=Contract.Kind.CLIENT_FORWARDING,
            expeditor=self.company_profile,
            customer=self.customer,
            contract_date=date.today(),
            status=Contract.Status.TERMINATED,
        )
        with self.assertRaises(ValidationError) as error:
            contract.full_clean()
        self.assertIn("terminated_on", error.exception.message_dict)

    def test_authenticated_user_can_create_contract(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("contract-create"),
            {
                "kind": Contract.Kind.CLIENT_FORWARDING,
                "expeditor": self.company_profile.pk,
                "customer": self.customer.pk,
                "number": "ТЭО-2026-001",
                "contract_date": "2026-08-28",
                "city": "Москва",
                "valid_until": "2027-08-28",
                "terminated_on": "",
                "status": Contract.Status.READY,
                "expeditor_representative": self.company_profile.director_name,
                "expeditor_authority_basis": "Устава",
                "counterparty_representative": self.customer.director_name,
                "counterparty_authority_basis": "Устава",
                "payment_term_days": "10",
                "payment_terms": "Оплата после получения оригиналов документов.",
                "debt_limit": "250000.00",
                "notes": "Проверен юридическим отделом",
            },
        )
        contract = Contract.objects.get(number="ТЭО-2026-001")
        self.assertRedirects(response, contract.get_absolute_url())
        self.assertEqual(contract.created_by, self.user)
        self.assertEqual(contract.customer, self.customer)
        self.assertIsNone(contract.carrier)
        self.assertEqual(contract.payment_term_days, 10)
        self.assertEqual(contract.debt_limit, Decimal("250000.00"))
        self.assertIn("оригиналов", contract.payment_terms)
        detail = self.client.get(contract.get_absolute_url())
        self.assertContains(detail, "Лимит задолженности")
        self.assertContains(detail, "Условия оплаты")

    def test_contract_download_fills_each_supplied_template(self):
        variants = (
            (Contract.Kind.CLIENT_FORWARDING, self.customer, "КЛИЕНТ"),
            (Contract.Kind.CARRIER_TRANSPORT, self.carrier, "ПЕРЕВОЗЧИК"),
            (
                Contract.Kind.SUBCONTRACTOR_FORWARDING,
                self.carrier,
                "ЭКСПЕДИТОР",
            ),
        )
        self.client.force_login(self.user)
        for index, (kind, counterparty, role) in enumerate(variants, start=1):
            with self.subTest(kind=kind):
                contract = Contract.objects.create(
                    kind=kind,
                    number=f"ДОГ-{index}",
                    contract_date=date(2026, 8, 28),
                    city="Москва",
                    valid_until=date(2027, 12, 31),
                    expeditor=self.company_profile,
                    customer=(
                        counterparty
                        if kind == Contract.Kind.CLIENT_FORWARDING
                        else None
                    ),
                    carrier=(
                        counterparty
                        if kind != Contract.Kind.CLIENT_FORWARDING
                        else None
                    ),
                    expeditor_representative=self.company_profile.director_name,
                    counterparty_representative=counterparty.director_name,
                    created_by=self.user,
                )
                response = self.client.get(
                    reverse("contract-download", args=[contract.pk])
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    response["Content-Type"],
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                document = Document(BytesIO(b"".join(response.streaming_content)))
                text_content = "\n".join(
                    [paragraph.text for paragraph in document.paragraphs]
                    + [
                        cell.text
                        for table in document.tables
                        for row in table.rows
                        for cell in row.cells
                    ]
                )
                self.assertIn(f"№ {contract.number}", text_content)
                self.assertIn("«28» августа 2026 г.", text_content)
                self.assertIn("31.12.2027", text_content)
                self.assertIn(self.company_profile.name, text_content)
                self.assertIn(counterparty.name, text_content)
                self.assertIn(counterparty.settlement_account, text_content)
                self.assertIn(role, text_content)

    def test_contract_registry_filters_by_expeditor_and_kind(self):
        second = self.create_second_expeditor()
        first_contract = Contract.objects.create(
            kind=Contract.Kind.CLIENT_FORWARDING,
            number="CLIENT-FIRST",
            expeditor=self.company_profile,
            customer=self.customer,
        )
        Contract.objects.create(
            kind=Contract.Kind.CARRIER_TRANSPORT,
            number="CARRIER-SECOND",
            expeditor=second,
            carrier=self.carrier,
        )
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("contract-list"),
            {
                "expeditor": self.company_profile.pk,
                "kind": Contract.Kind.CLIENT_FORWARDING,
            },
        )
        self.assertContains(response, first_contract.number)
        self.assertNotContains(response, "CARRIER-SECOND")
        self.assertEqual(response.context["contract_count"], 1)

    def test_user_can_start_direct_chat(self):
        colleague = get_user_model().objects.create_user(
            username="colleague", password="test-password",
            first_name="Пётр", last_name="Смирнов",
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("chat-start"), {"user_id": colleague.pk}
        )
        conversation = DirectConversation.objects.get()
        self.assertRedirects(response, conversation.get_absolute_url())
        self.assertTrue(conversation.includes(self.user))
        self.assertTrue(conversation.includes(colleague))

    def test_user_can_send_chat_message(self):
        colleague = get_user_model().objects.create_user(
            username="sender-test", password="test-password"
        )
        conversation, _ = DirectConversation.get_or_create_between(
            self.user, colleague
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("chat-send", args=[conversation.pk]),
            {"text": "Статус заявки обновлён."},
        )
        self.assertRedirects(response, conversation.get_absolute_url())
        message = ChatMessage.objects.get()
        self.assertEqual(message.sender, self.user)
        self.assertEqual(message.text, "Статус заявки обновлён.")

    def test_ajax_chat_send_returns_message_payload(self):
        colleague = get_user_model().objects.create_user(
            username="ajax-colleague", password="test-password"
        )
        conversation, _ = DirectConversation.get_or_create_between(
            self.user, colleague
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("chat-send", args=[conversation.pk]),
            {"text": "Сообщение без перезагрузки"},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["message"]["is_mine"])
        self.assertEqual(
            response.json()["message"]["text"], "Сообщение без перезагрузки"
        )

    def test_chat_conversation_is_private_to_participants(self):
        first = get_user_model().objects.create_user(
            username="private-first", password="test-password"
        )
        second = get_user_model().objects.create_user(
            username="private-second", password="test-password"
        )
        conversation, _ = DirectConversation.get_or_create_between(first, second)
        self.client.force_login(self.user)
        self.assertEqual(
            self.client.get(conversation.get_absolute_url()).status_code, 404
        )
        self.assertEqual(
            self.client.post(
                reverse("chat-send", args=[conversation.pk]), {"text": "Нет доступа"}
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.get(
                reverse("chat-messages", args=[conversation.pk])
            ).status_code,
            404,
        )

    def test_unread_chat_badge_and_read_receipt(self):
        colleague = get_user_model().objects.create_user(
            username="unread-colleague", password="test-password"
        )
        conversation, _ = DirectConversation.get_or_create_between(
            self.user, colleague
        )
        message = ChatMessage.objects.create(
            conversation=conversation,
            sender=colleague,
            text="Непрочитанное сообщение",
        )
        self.client.force_login(self.user)
        dashboard = self.client.get(reverse("dashboard"))
        self.assertEqual(dashboard.context["chat_unread_count"], 1)
        self.assertContains(dashboard, "nav-unread")
        chat = self.client.get(conversation.get_absolute_url())
        self.assertContains(chat, "Непрочитанное сообщение")
        message.refresh_from_db()
        self.assertIsNotNone(message.read_at)

    def test_chat_polling_returns_only_new_messages(self):
        colleague = get_user_model().objects.create_user(
            username="poll-colleague", password="test-password",
            first_name="Ольга",
        )
        conversation, _ = DirectConversation.get_or_create_between(
            self.user, colleague
        )
        first = ChatMessage.objects.create(
            conversation=conversation, sender=self.user, text="Первое"
        )
        second = ChatMessage.objects.create(
            conversation=conversation, sender=colleague, text="Второе"
        )
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("chat-messages", args=[conversation.pk]),
            {"after": first.pk},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["messages"]), 1)
        self.assertEqual(response.json()["messages"][0]["id"], second.pk)
        self.assertFalse(response.json()["messages"][0]["is_mine"])

    def test_chat_message_author_can_edit_and_delete_own_message(self):
        colleague = get_user_model().objects.create_user(
            username="message-actions", password="test-password"
        )
        conversation, _ = DirectConversation.get_or_create_between(
            self.user, colleague
        )
        own_message = ChatMessage.objects.create(
            conversation=conversation,
            sender=self.user,
            text="Исходный текст",
        )
        foreign_message = ChatMessage.objects.create(
            conversation=conversation,
            sender=colleague,
            text="Чужой текст",
        )
        self.client.force_login(self.user)

        chat = self.client.get(conversation.get_absolute_url())
        self.assertContains(chat, own_message.get_absolute_url())
        self.assertContains(chat, own_message.get_delete_url())
        self.assertNotContains(chat, foreign_message.get_absolute_url())

        response = self.client.post(
            own_message.get_absolute_url(),
            {"text": "Исправленный текст"},
        )
        self.assertRedirects(response, conversation.get_absolute_url())
        own_message.refresh_from_db()
        self.assertEqual(own_message.text, "Исправленный текст")
        self.assertIsNotNone(own_message.edited_at)

        self.assertEqual(
            self.client.post(foreign_message.get_absolute_url(), {"text": "Взлом"}).status_code,
            404,
        )
        self.assertEqual(self.client.post(foreign_message.get_delete_url()).status_code, 404)
        self.assertTrue(ChatMessage.objects.filter(pk=foreign_message.pk).exists())

        response = self.client.post(own_message.get_delete_url())
        self.assertRedirects(response, conversation.get_absolute_url())
        self.assertFalse(ChatMessage.objects.filter(pk=own_message.pk).exists())

    def test_chat_users_can_exchange_and_securely_download_files(self):
        colleague = get_user_model().objects.create_user(
            username="file-colleague", password="test-password"
        )
        outsider = get_user_model().objects.create_user(
            username="file-outsider", password="test-password"
        )
        conversation, _ = DirectConversation.get_or_create_between(
            self.user, colleague
        )
        upload = SimpleUploadedFile(
            "route_sheet.pdf",
            b"%PDF-1.4 test chat attachment",
            content_type="application/pdf",
        )

        with tempfile.TemporaryDirectory() as media_directory:
            with override_settings(MEDIA_ROOT=media_directory):
                self.client.force_login(self.user)
                response = self.client.post(
                    reverse("chat-send", args=[conversation.pk]),
                    {"text": "", "attachment": upload},
                    HTTP_X_REQUESTED_WITH="XMLHttpRequest",
                )
                self.assertEqual(response.status_code, 200)
                payload = response.json()["message"]
                self.assertEqual(payload["attachment"]["name"], "route_sheet.pdf")
                message = ChatMessage.objects.get()
                self.assertEqual(message.attachment_name, "route_sheet.pdf")
                self.assertGreater(message.attachment_size, 0)

                self.client.force_login(colleague)
                download = self.client.get(message.get_attachment_url())
                self.assertEqual(download.status_code, 200)
                self.assertEqual(
                    b"".join(download.streaming_content),
                    b"%PDF-1.4 test chat attachment",
                )

                self.client.force_login(outsider)
                self.assertEqual(
                    self.client.get(message.get_attachment_url()).status_code,
                    404,
                )

    def test_chat_rejects_executable_attachments(self):
        colleague = get_user_model().objects.create_user(
            username="unsafe-file", password="test-password"
        )
        conversation, _ = DirectConversation.get_or_create_between(
            self.user, colleague
        )
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("chat-send", args=[conversation.pk]),
            {
                "text": "",
                "attachment": SimpleUploadedFile(
                    "dangerous.exe", b"not an executable", content_type="application/octet-stream"
                ),
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(ChatMessage.objects.exists())

    def test_chat_full_sync_reflects_edits_and_deletions(self):
        colleague = get_user_model().objects.create_user(
            username="sync-actions", password="test-password"
        )
        conversation, _ = DirectConversation.get_or_create_between(
            self.user, colleague
        )
        message = ChatMessage.objects.create(
            conversation=conversation,
            sender=self.user,
            text="До изменения",
        )
        self.client.force_login(self.user)
        self.client.post(message.get_absolute_url(), {"text": "После изменения"})

        response = self.client.get(
            reverse("chat-messages", args=[conversation.pk]),
            {"sync": "1"},
        )
        self.assertTrue(response.json()["full_sync"])
        self.assertEqual(response.json()["messages"][0]["text"], "После изменения")
        self.assertTrue(response.json()["messages"][0]["edited"])

        self.client.post(message.get_delete_url())
        response = self.client.get(
            reverse("chat-messages", args=[conversation.pk]),
            {"sync": "1"},
        )
        self.assertEqual(response.json()["messages"], [])

    def test_dashboard_metric_cards_link_to_filtered_lists(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"))
        for scope in (
            "active", "unassigned", "revenue", "margin", "receivables", "payables"
        ):
            with self.subTest(scope=scope):
                self.assertContains(
                    response, f'{reverse("shipment-list")}?scope={scope}'
                )

    def test_dashboard_scopes_filter_shipments(self):
        common = {
            "expeditor": self.company_profile,
            "customer": self.customer,
            "manager": self.user,
            "cargo_name": "Тестовый груз",
            "pickup_city": "Москва",
            "pickup_date": date.today(),
            "delivery_city": "Тула",
            "delivery_date": date.today() + timedelta(days=1),
        }
        unassigned = Shipment.objects.create(
            number="SCOPE-UNASSIGNED",
            status=Shipment.Status.NEW,
            customer_price=Decimal("50000"),
            carrier_price=Decimal("0"),
            **common,
        )
        loss = Shipment.objects.create(
            number="SCOPE-LOSS",
            carrier=self.carrier,
            status=Shipment.Status.CLOSED,
            customer_price=Decimal("50000"),
            carrier_price=Decimal("70000"),
            **common,
        )
        cancelled = Shipment.objects.create(
            number="SCOPE-CANCELLED",
            carrier=self.carrier,
            status=Shipment.Status.CANCELLED,
            customer_price=Decimal("90000"),
            carrier_price=Decimal("10000"),
            **common,
        )
        expected_ids = {
            "active": {self.shipment.pk, unassigned.pk},
            "unassigned": {unassigned.pk},
            "revenue": {self.shipment.pk, unassigned.pk, loss.pk},
            "margin": {self.shipment.pk, unassigned.pk},
            "receivables": {self.shipment.pk, unassigned.pk, loss.pk},
            "payables": {self.shipment.pk, loss.pk},
        }
        self.client.force_login(self.user)
        for scope, expected in expected_ids.items():
            with self.subTest(scope=scope):
                response = self.client.get(
                    reverse("shipment-list"), {"scope": scope}
                )
                actual = {shipment.pk for shipment in response.context["shipments"]}
                self.assertEqual(actual, expected)
                self.assertNotIn(cancelled.pk, actual)

    def test_dashboard_shows_current_debt_balances(self):
        Payment.objects.create(
            shipment=self.shipment,
            direction=Payment.Direction.INCOME,
            amount=Decimal("40000"),
            payment_date=date.today(),
            created_by=self.user,
        )
        Payment.objects.create(
            shipment=self.shipment,
            direction=Payment.Direction.EXPENSE,
            amount=Decimal("20000"),
            payment_date=date.today(),
            created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"), {"currency": "RUB"})
        self.assertEqual(response.context["receivables"], Decimal("60000"))
        self.assertEqual(response.context["payables"], Decimal("55000"))
        self.assertContains(response, "Дебиторская задолженность")
        self.assertContains(response, "Кредиторская задолженность")

    def test_dashboard_financial_cards_filter_by_currency(self):
        second = self.create_second_expeditor()
        usd_shipment = self.create_shipment_for(second, number="USD-001")
        usd_shipment.currency = "USD"
        usd_shipment.customer_price = Decimal("5000")
        usd_shipment.carrier_price = Decimal("3500")
        usd_shipment.save()
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"), {"currency": "USD"})
        self.assertEqual(response.context["revenue"], Decimal("5000"))
        self.assertEqual(response.context["margin"], Decimal("1500"))
        self.assertEqual(response.context["receivables"], Decimal("5000"))
        self.assertEqual(response.context["payables"], Decimal("3500"))

    def test_shipment_search(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("shipment-list"), {"q": "Казань"})
        self.assertContains(response, self.shipment.number)
        response = self.client.get(reverse("shipment-list"), {"q": "Владивосток"})
        self.assertNotContains(response, self.shipment.number)

    def test_delivery_cannot_precede_pickup(self):
        form = ShipmentForm(
            data={
                "expeditor": self.company_profile.pk,
                "customer": self.customer.pk,
                "manager": self.user.pk,
                "status": Shipment.Status.NEW,
                "cargo_name": "Груз",
                "weight_kg": "100",
                "volume_m3": "1",
                "pickup_city": "Москва",
                "pickup_date": "2026-08-20",
                "delivery_city": "Тула",
                "delivery_date": "2026-08-19",
                "customer_price": "1000",
                "carrier_price": "800",
                "currency": "RUB",
                "payment_status": Shipment.PaymentStatus.AWAITING,
            }
        )
        self.assertFalse(form.is_valid())
        self.assertIn("delivery_date", form.errors)

    def test_edit_form_renders_dates_in_russian_text_format(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("shipment-update", args=[self.shipment.pk]))
        self.assertContains(
            response,
            'value="{}"'.format(self.shipment.pickup_date.strftime("%d.%m.%Y")),
        )
        self.assertContains(
            response,
            'value="{}"'.format(self.shipment.delivery_date.strftime("%d.%m.%Y")),
        )
        self.assertContains(response, 'placeholder="ДД.ММ.ГГГГ"')
        self.assertContains(response, "data-crm-date")
        self.assertNotContains(response, 'type="date"')

    def test_authenticated_user_can_create_customer(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("customer-create"),
            {"name": "Новый клиент", "is_active": "on"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Customer.objects.filter(name="Новый клиент").exists())

    def test_payments_calculate_both_debt_balances(self):
        self.shipment.customer_payment_due_date = date.today() - timedelta(days=1)
        self.shipment.carrier_payment_due_date = date.today() - timedelta(days=1)
        self.shipment.save()
        Payment.objects.create(
            shipment=self.shipment,
            direction=Payment.Direction.INCOME,
            amount=Decimal("40000"),
            payment_date=date.today(),
            created_by=self.user,
        )
        Payment.objects.create(
            shipment=self.shipment,
            direction=Payment.Direction.EXPENSE,
            amount=Decimal("25000"),
            payment_date=date.today(),
            created_by=self.user,
        )
        self.assertEqual(self.shipment.received_amount, Decimal("40000"))
        self.assertEqual(self.shipment.paid_to_carrier_amount, Decimal("25000"))
        self.assertEqual(self.shipment.receivable_balance, Decimal("60000"))
        self.assertEqual(self.shipment.payable_balance, Decimal("50000"))
        self.assertEqual(
            self.shipment.receivable_state, Shipment.PaymentStatus.OVERDUE
        )
        self.assertEqual(self.shipment.payable_state, Shipment.PaymentStatus.OVERDUE)

    def test_full_customer_payment_marks_receivable_as_paid(self):
        Payment.objects.create(
            shipment=self.shipment,
            direction=Payment.Direction.INCOME,
            amount=self.shipment.customer_price,
            payment_date=date.today(),
            created_by=self.user,
        )
        self.shipment.refresh_from_db()
        self.assertEqual(self.shipment.receivable_balance, Decimal("0"))
        self.assertEqual(self.shipment.receivable_state, Shipment.PaymentStatus.PAID)
        self.assertEqual(self.shipment.payment_status, Shipment.PaymentStatus.PAID)

    def test_carrier_payment_requires_assigned_carrier(self):
        self.shipment.carrier = None
        self.shipment.save()
        payment = Payment(
            shipment=self.shipment,
            direction=Payment.Direction.EXPENSE,
            amount=Decimal("1000"),
            payment_date=date.today(),
        )
        with self.assertRaises(ValidationError) as error:
            payment.full_clean()
        self.assertIn("direction", error.exception.message_dict)

    def test_planner_contains_transportation_events_and_task_filters(self):
        transportation = self.shipment.transportation
        transportation.planned_start_date = date.today()
        transportation.planned_end_date = date.today() + timedelta(days=1)
        transportation.customer_payment_due_date = date.today() + timedelta(days=3)
        transportation.executor_payment_due_date = date.today() + timedelta(days=4)
        transportation.save(
            update_fields=[
                "planned_start_date", "planned_end_date",
                "customer_payment_due_date", "executor_payment_due_date",
                "updated_at",
            ]
        )
        task = PlannerTask.objects.create(
            transportation=transportation,
            title="Запросить подписанный УПД",
            kind=PlannerTask.Kind.DOCUMENT,
            priority=PlannerTask.Priority.HIGH,
            assignee=self.user,
            due_date=date.today(),
        )
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("planner"),
            {
                "date_from": date.today().strftime("%d.%m.%Y"),
                "date_to": (date.today() + timedelta(days=5)).strftime("%d.%m.%Y"),
                "assignee": str(self.user.pk),
                "task_status": PlannerTask.Status.TODO,
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["event_count"], 6)
        self.assertEqual(response.context["tasks"], [task])
        self.assertContains(response, "Планировщик")
        self.assertContains(response, "Погрузка")
        self.assertContains(response, "Запросить подписанный УПД")
        self.assertContains(response, reverse("planner-task-update", args=[task.pk]))

    def test_planner_task_can_be_completed_from_agenda(self):
        task = PlannerTask.objects.create(
            title="Проверить документы",
            assignee=self.user,
            due_date=date.today(),
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("planner-task-complete", args=[task.pk]),
            {"next": reverse("planner")},
        )
        self.assertRedirects(response, reverse("planner"))
        task.refresh_from_db()
        self.assertEqual(task.status, PlannerTask.Status.DONE)
        self.assertEqual(task.completed_by, self.user)
        self.assertIsNotNone(task.completed_at)

    def test_planner_task_form_accepts_crm_date_format(self):
        self.client.force_login(self.user)
        due_date = date.today() + timedelta(days=2)
        response = self.client.post(
            reverse("planner-task-create"),
            {
                "title": "Подтвердить время выгрузки",
                "description": "Позвонить получателю.",
                "kind": PlannerTask.Kind.CONTROL,
                "status": PlannerTask.Status.TODO,
                "priority": PlannerTask.Priority.NORMAL,
                "transportation": self.shipment.transportation.pk,
                "assignee": self.user.pk,
                "due_date": due_date.strftime("%d.%m.%Y"),
            },
        )
        self.assertRedirects(response, reverse("planner"))
        task = PlannerTask.objects.get(title="Подтвердить время выгрузки")
        self.assertEqual(task.due_date, due_date)
        self.assertEqual(task.transportation_id, self.shipment.transportation.pk)

    def test_planner_builds_automatic_control_tasks_from_trip_state(self):
        transportation = self.shipment.transportation
        transportation.planned_start_date = date.today()
        transportation.save(update_fields=["planned_start_date", "updated_at"])
        self.client.force_login(self.user)

        planner = self.client.get(reverse("planner"))
        titles = {item["title"] for item in planner.context["automatic_tasks"]}

        self.assertIn("Создать договор с исполнителем", titles)
        self.assertIn("Назначить водителя", titles)
        self.assertIn("Назначить транспорт", titles)
        self.assertContains(planner, "Уведомления по рейсам")

        dashboard = self.client.get(reverse("dashboard"))
        self.assertGreater(dashboard.context["dashboard_notification_count"], 0)
        self.assertContains(dashboard, "Требуют внимания")

    def test_reports_show_financial_totals_and_debts(self):
        Payment.objects.create(
            shipment=self.shipment,
            direction=Payment.Direction.INCOME,
            amount=Decimal("40000"),
            payment_date=date.today(),
            created_by=self.user,
        )
        Payment.objects.create(
            shipment=self.shipment,
            direction=Payment.Direction.EXPENSE,
            amount=Decimal("20000"),
            payment_date=date.today(),
            created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse("reports"), {"currency": "RUB"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["totals"]["revenue"], Decimal("100000"))
        self.assertEqual(response.context["totals"]["margin"], Decimal("25000"))
        self.assertEqual(response.context["totals"]["received"], Decimal("40000"))
        self.assertEqual(
            response.context["totals"]["receivables"], Decimal("60000")
        )
        self.assertEqual(response.context["totals"]["payables"], Decimal("55000"))
        self.assertEqual(response.context["tax_transportation_count"], 1)
        self.assertEqual(
            response.context["tax_totals"]["sales_vat"], Decimal("18032.79")
        )
        self.assertEqual(
            response.context["tax_totals"]["vat_payable"], Decimal("18032.79")
        )
        self.assertEqual(
            response.context["tax_totals"]["profit"], Decimal("6967.21")
        )
        self.assertEqual(
            response.context["tax_totals"]["profit_tax"], Decimal("1741.80")
        )
        self.assertEqual(
            response.context["tax_totals"]["net_profit"], Decimal("5225.41")
        )
        self.assertContains(response, "Дебиторская задолженность")
        self.assertContains(response, "Кредиторская задолженность")
        self.assertContains(response, "Сводный отчёт по налогам")
        self.assertContains(response, "Показать расшифровку по рейсам")

    def test_reports_include_trip_analytics_and_role_filters(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("reports"), {"currency": "RUB"})
        self.assertEqual(response.context["trip_totals"]["trip_count"], 1)
        self.assertEqual(
            response.context["trip_totals"]["revenue"], Decimal("100000")
        )
        self.assertEqual(
            response.context["trip_totals"]["margin"], Decimal("25000")
        )
        self.assertContains(response, "Аналитика по рейсам")
        self.assertContains(response, "Возраст задолженности по рейсам")
        self.assertContains(response, "Показать аналитику по участникам рейсов")

        transportation = self.shipment.transportation
        filtered = self.client.get(
            reverse("reports"),
            {
                "currency": "RUB",
                "trip_status": transportation.status,
                "trip_client": self.customer.organization.pk,
            },
        )
        self.assertEqual(filtered.context["trip_totals"]["trip_count"], 1)
        empty = self.client.get(
            reverse("reports"),
            {"currency": "RUB", "trip_status": Transportation.Status.CLOSED},
        )
        self.assertEqual(empty.context["trip_totals"]["trip_count"], 0)

    def test_debt_report_shows_open_trip_balances(self):
        self.shipment.driver = self.driver
        self.shipment.vehicle = self.vehicle
        self.shipment.save()
        transportation = self.shipment.transportation
        assignment = transportation.active_vehicle_assignment()
        assignment.trailer_registration_number = "В456ВВ198"
        assignment.save(update_fields=["trailer_registration_number", "updated_at"])
        contract = Contract.objects.create(
            kind=Contract.Kind.CARRIER_TRANSPORT,
            expeditor=self.company_profile,
            carrier=self.carrier,
            contract_date=date.today(),
        )
        link = transportation.active_execution_link()
        link.contract = contract
        link.save(update_fields=["contract", "updated_at"])
        self.client.force_login(self.user)
        self.client.post(reverse("transportation-post", args=[transportation.pk]))

        response = self.client.get(reverse("debt-report"), {"currency": "RUB"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["debt_totals"]["receivable"], Decimal("100000"))
        self.assertEqual(response.context["debt_totals"]["payable"], Decimal("75000"))
        self.assertContains(response, "Задолженность по рейсам")
        self.assertContains(response, "Тестовый клиент")
        self.assertContains(response, "Тестовый перевозчик")
        self.assertContains(
            response,
            reverse("transportation-payment-create", args=[transportation.pk]),
        )

        receivable_only = self.client.get(
            reverse("debt-report"),
            {"currency": "RUB", "side": SettlementMovement.Side.RECEIVABLE},
        )
        self.assertEqual(len(receivable_only.context["debt_rows"]), 1)
        self.assertEqual(
            receivable_only.context["debt_rows"][0]["side"],
            SettlementMovement.Side.RECEIVABLE,
        )

    def test_profitability_report_shows_trip_financials_and_filters(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("profitability-report"), {"currency": "RUB"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["profit_totals"]["trip_count"], 1)
        self.assertEqual(response.context["profit_totals"]["revenue"], Decimal("100000"))
        self.assertEqual(response.context["profit_totals"]["cost"], Decimal("75000"))
        self.assertEqual(response.context["profit_totals"]["margin"], Decimal("25000"))
        self.assertContains(response, "Прибыльность рейсов")
        self.assertContains(response, "Тестовый клиент")
        self.assertContains(response, "Тестовый перевозчик")
        self.assertContains(response, reverse("transportation-detail", args=[self.shipment.transportation.pk]))

        low_margin = self.client.get(
            reverse("profitability-report"),
            {"currency": "RUB", "low_margin_percent": "10"},
        )
        self.assertEqual(low_margin.context["profit_totals"]["trip_count"], 0)
        loss_only = self.client.get(
            reverse("profitability-report"),
            {"currency": "RUB", "loss_only": "1"},
        )
        self.assertEqual(loss_only.context["profit_totals"]["trip_count"], 0)

    def test_financial_reports_export_xlsx(self):
        self.client.force_login(self.user)
        for url_name, expected_sheet in (
            ("debt-report-export", "Рейсы"),
            ("profitability-report-export", "Клиенты"),
            ("tax-report-export", "Компании"),
        ):
            with self.subTest(url_name=url_name):
                response = self.client.get(reverse(url_name), {"currency": "RUB"})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    response["Content-Type"],
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                with ZipFile(BytesIO(response.content)) as archive:
                    workbook = archive.read("xl/workbook.xml").decode("utf-8")
                    self.assertIn(expected_sheet, workbook)
                    self.assertIn("xl/worksheets/sheet1.xml", archive.namelist())

    def test_reports_filter_by_pickup_period(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("reports"),
            {"date_from": (date.today() + timedelta(days=1)).isoformat()},
        )
        self.assertEqual(response.context["shipment_count"], 0)
        self.assertEqual(response.context["totals"]["revenue"], Decimal("0"))
        self.assertEqual(response.context["tax_transportation_count"], 0)
        self.assertEqual(response.context["tax_totals"]["vat_payable"], Decimal("0"))

    def test_reports_filter_accepts_russian_date_format(self):
        self.client.force_login(self.user)
        filter_date = date.today() + timedelta(days=1)
        response = self.client.get(
            reverse("reports"),
            {"date_from": filter_date.strftime("%d.%m.%Y")},
        )
        self.assertEqual(response.context["shipment_count"], 0)
        self.assertEqual(
            response.context["current_date_from"],
            filter_date.strftime("%d.%m.%Y"),
        )
        self.assertContains(response, 'placeholder="ДД.ММ.ГГГГ"')
        self.assertNotContains(response, 'type="date"')

    def test_shipment_payment_filter_uses_current_due_date(self):
        self.shipment.customer_payment_due_date = date.today() - timedelta(days=1)
        self.shipment.payment_status = Shipment.PaymentStatus.AWAITING
        self.shipment.save()
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("shipment-list"),
            {"payment": Shipment.PaymentStatus.OVERDUE},
        )
        self.assertContains(response, self.shipment.number)

    def test_authenticated_user_can_register_payment(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("payment-create", args=[self.shipment.pk]),
            {
                "direction": Payment.Direction.INCOME,
                "amount": "15000.00",
                "payment_date": date.today().isoformat(),
                "method": Payment.Method.BANK,
                "reference": "ПП-101",
                "notes": "Аванс",
            },
        )
        self.assertRedirects(response, self.shipment.get_absolute_url())
        payment = Payment.objects.get(reference="ПП-101")
        self.assertEqual(payment.amount, Decimal("15000.00"))
        self.assertEqual(payment.created_by, self.user)

    def test_forwarding_order_form_prefills_customer_and_shipment_data(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("forwarding-order", args=[self.shipment.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.customer.name)
        self.assertContains(response, self.shipment.cargo_name)
        self.assertContains(response, self.shipment.pickup_city)
        self.assertContains(response, "Скачать DOCX")

    def test_forwarding_order_download_creates_valid_docx(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("forwarding-order", args=[self.shipment.pk]),
            {
                "action": "download",
                "contract_number": "ТЭ-55/2026",
                "contract_date": "2026-08-01",
                "order_date": "2026-08-28",
                "shipper_name": self.customer.name,
                "shipper_tax_id": "7700000000",
                "shipper_address": "Москва, ул. Тестовая, 1",
                "shipper_contact_name": "Иван Иванов",
                "shipper_phone": "+7 900 100-20-30",
                "consignee_name": "ООО Получатель",
                "consignee_address": "Казань, ул. Складская, 2",
                "consignee_contact_name": "Пётр Петров",
                "consignee_phone": "+7 900 300-40-50",
                "pickup_hours": "09:00–18:00",
                "packaging_type": "Паллеты",
                "package_count": "10",
                "cargo_length_m": "2.40",
                "cargo_width_m": "1.20",
                "cargo_height_m": "1.50",
                "special_conditions": "Требуется крепление",
                "cargo_insurance": ForwardingOrder.Insurance.YES,
                "payer": self.customer.name,
                "payment_place": "Безналичный расчёт",
                "expediter_representative": "Анна Петрова",
                "client_representative": "Иван Иванов",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("attachment", response["Content-Disposition"])
        content = b"".join(response.streaming_content)
        document = Document(BytesIO(content))
        document_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
        )
        self.assertIn("ПОРУЧЕНИЕ ЭКСПЕДИТОРУ", document_text)
        self.assertIn("ТЭ-55/2026", document_text)
        self.assertIn(self.shipment.cargo_name, document_text)
        self.assertIn(self.shipment.pickup_city, document_text)
        self.assertIn("ООО Получатель", document_text)
        self.assertTrue(
            ForwardingOrder.objects.filter(shipment=self.shipment).exists()
        )

    def test_forwarding_order_draft_is_saved(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("forwarding-order", args=[self.shipment.pk]),
            {
                "action": "save",
                "contract_number": "DRAFT-01",
                "cargo_insurance": ForwardingOrder.Insurance.NOT_SPECIFIED,
            },
        )
        self.assertRedirects(response, self.shipment.get_absolute_url())
        self.assertEqual(
            ForwardingOrder.objects.get(shipment=self.shipment).contract_number,
            "DRAFT-01",
        )

    def test_authenticated_user_can_create_expected_shipment_document(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("shipment-document-create", args=[self.shipment.pk]),
            {
                "kind": ShipmentDocument.Kind.TRANSPORT_WAYBILL,
                "party": ShipmentDocument.Party.CARRIER,
                "status": ShipmentDocument.Status.EXPECTED,
                "number": "",
                "document_date": "",
                "expected_date": (date.today() - timedelta(days=1)).isoformat(),
                "amount": "",
                "currency": "RUB",
                "notes": "Ждём оригинал от перевозчика",
            },
        )
        self.assertRedirects(response, self.shipment.get_absolute_url())
        document_record = ShipmentDocument.objects.get(shipment=self.shipment)
        self.assertEqual(
            document_record.kind, ShipmentDocument.Kind.TRANSPORT_WAYBILL
        )
        self.assertTrue(document_record.is_overdue)
        self.assertEqual(document_record.created_by, self.user)

    def test_document_can_be_created_from_general_registry(self):
        carrier_organization = Organization.objects.create(
            name='ООО "Документный перевозчик"',
            tax_id="7722555000",
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("shipment-document-create-general"),
            {
                "shipment": self.shipment.pk,
                "direction": ShipmentDocument.Direction.INCOMING,
                "counterparty": carrier_organization.pk,
                "kind": ShipmentDocument.Kind.INVOICE,
                "party": ShipmentDocument.Party.CARRIER,
                "status": ShipmentDocument.Status.RECEIVED,
                "number": "ВХ-001",
                "document_date": date.today().isoformat(),
                "expected_date": "",
                "amount": "40000.00",
                "vat_amount": "0.00",
                "currency": "RUB",
                "notes": "Входящий счёт перевозчика",
            },
        )
        self.assertRedirects(response, reverse("shipment-document-list"))
        document_record = ShipmentDocument.objects.get(number="ВХ-001")
        self.assertEqual(document_record.shipment, self.shipment)
        self.assertEqual(document_record.direction, ShipmentDocument.Direction.INCOMING)
        self.assertEqual(document_record.counterparty, carrier_organization)
        self.assertEqual(document_record.vat_amount, Decimal("0.00"))

    def test_document_registry_filters_by_kind_and_status(self):
        ShipmentDocument.objects.create(
            shipment=self.shipment,
            kind=ShipmentDocument.Kind.UPD,
            party=ShipmentDocument.Party.CUSTOMER,
            status=ShipmentDocument.Status.SIGNED,
            number="УПД-001",
            created_by=self.user,
        )
        ShipmentDocument.objects.create(
            shipment=self.shipment,
            kind=ShipmentDocument.Kind.INVOICE,
            party=ShipmentDocument.Party.CUSTOMER,
            status=ShipmentDocument.Status.ISSUED,
            number="СЧ-001",
            created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("shipment-document-list"),
            {
                "kind": ShipmentDocument.Kind.UPD,
                "status": ShipmentDocument.Status.SIGNED,
            },
        )
        self.assertContains(response, "УПД-001")
        self.assertNotContains(response, "СЧ-001")

    def test_document_registry_filters_by_direction_and_counterparty(self):
        carrier_organization = Organization.objects.create(
            name='ООО "Реестровый перевозчик"',
            tax_id="7722666000",
        )
        customer_organization = Organization.objects.create(
            name='ООО "Реестровый клиент"',
            tax_id="7711666000",
        )
        ShipmentDocument.objects.create(
            shipment=self.shipment,
            direction=ShipmentDocument.Direction.INCOMING,
            counterparty=carrier_organization,
            kind=ShipmentDocument.Kind.INVOICE,
            party=ShipmentDocument.Party.CARRIER,
            status=ShipmentDocument.Status.RECEIVED,
            number="ВХ-ПЕРЕВОЗЧИК",
            created_by=self.user,
        )
        ShipmentDocument.objects.create(
            shipment=self.shipment,
            direction=ShipmentDocument.Direction.OUTGOING,
            counterparty=customer_organization,
            kind=ShipmentDocument.Kind.UPD,
            party=ShipmentDocument.Party.CUSTOMER,
            status=ShipmentDocument.Status.ISSUED,
            number="ИСХ-КЛИЕНТ",
            created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("shipment-document-list"),
            {"direction": ShipmentDocument.Direction.INCOMING, "q": carrier_organization.tax_id},
        )
        self.assertContains(response, "ВХ-ПЕРЕВОЗЧИК")
        self.assertNotContains(response, "ИСХ-КЛИЕНТ")

    def test_shipment_document_can_be_deleted(self):
        customer_organization = Organization.objects.create(
            name='ООО "Документный клиент"',
            tax_id="7711777000",
        )
        document_record = ShipmentDocument.objects.create(
            shipment=self.shipment,
            direction=ShipmentDocument.Direction.OUTGOING,
            counterparty=customer_organization,
            kind=ShipmentDocument.Kind.ACT,
            party=ShipmentDocument.Party.CUSTOMER,
            status=ShipmentDocument.Status.DRAFT,
            number="АКТ-DEL",
            created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("shipment-document-delete", args=[document_record.pk])
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, self.shipment.get_absolute_url())
        self.assertFalse(
            ShipmentDocument.objects.filter(pk=document_record.pk).exists()
        )

    def test_transportation_document_update_redirects_to_transportation(self):
        transportation = self.shipment.transportation
        document_record = ShipmentDocument.objects.create(
            transportation=transportation,
            direction=ShipmentDocument.Direction.OUTGOING,
            kind=ShipmentDocument.Kind.UPD,
            party=ShipmentDocument.Party.CUSTOMER,
            status=ShipmentDocument.Status.DRAFT,
            number="УПД-РЕЙС",
            amount=Decimal("100000.00"),
            currency="RUB",
            created_by=self.user,
        )
        upload = SimpleUploadedFile("upd.xml", b"<xml></xml>", content_type="text/xml")
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("shipment-document-update", args=[document_record.pk]),
            {
                "shipment": "",
                "transportation": transportation.pk,
                "direction": ShipmentDocument.Direction.OUTGOING,
                "kind": ShipmentDocument.Kind.UPD,
                "party": ShipmentDocument.Party.CUSTOMER,
                "status": ShipmentDocument.Status.ISSUED,
                "number": "УПД-РЕЙС-1",
                "document_date": date.today().isoformat(),
                "expected_date": "",
                "amount": "100000.00",
                "vat_amount": "16666.67",
                "currency": "RUB",
                "notes": "Файл приложен",
                "file": upload,
            },
        )
        self.assertRedirects(response, transportation.get_absolute_url())
        document_record.refresh_from_db()
        self.assertEqual(document_record.number, "УПД-РЕЙС-1")
        self.assertTrue(document_record.file.name.startswith(f"transportations/{transportation.pk}/documents/"))

    def test_transportation_document_can_be_deleted(self):
        transportation = self.shipment.transportation
        document_record = ShipmentDocument.objects.create(
            transportation=transportation,
            direction=ShipmentDocument.Direction.INCOMING,
            kind=ShipmentDocument.Kind.ACT,
            party=ShipmentDocument.Party.CARRIER,
            status=ShipmentDocument.Status.RECEIVED,
            number="АКТ-РЕЙС-DEL",
            created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("shipment-document-delete", args=[document_record.pk])
        )
        self.assertRedirects(response, transportation.get_absolute_url())
        self.assertFalse(
            ShipmentDocument.objects.filter(pk=document_record.pk).exists()
        )

    def test_document_batch_posts_documents_for_selected_transportations(self):
        transportation = self.shipment.transportation
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("document-batch-create"),
            {
                "direction": DocumentBatch.Direction.OUTGOING,
                "document_date": date.today().isoformat(),
                "owner_company": transportation.owner_company_id,
                "default_kind": ShipmentDocument.Kind.UPD,
                "default_status": ShipmentDocument.Status.ISSUED,
                "currency": "RUB",
                "reference": "РЕЕСТР-001",
                "notes": "Массовая реализация",
                "transportation_ids": [transportation.pk],
                f"kind_{transportation.pk}": ShipmentDocument.Kind.UPD,
                f"number_{transportation.pk}": "УПД-ПАЧКА-001",
                f"date_{transportation.pk}": date.today().isoformat(),
                f"amount_{transportation.pk}": "100000.00",
                f"vat_{transportation.pk}": "18032.79",
                "action": "post",
            },
        )
        batch = DocumentBatch.objects.get(reference="РЕЕСТР-001")
        self.assertRedirects(response, batch.get_absolute_url())
        self.assertEqual(batch.status, DocumentBatch.Status.POSTED)
        line = DocumentBatchLine.objects.get(batch=batch)
        self.assertIsNotNone(line.shipment_document)
        self.assertEqual(line.shipment_document.shipment, self.shipment)
        self.assertEqual(line.shipment_document.number, "УПД-ПАЧКА-001")
        self.assertEqual(line.shipment_document.direction, ShipmentDocument.Direction.OUTGOING)
        self.assertEqual(line.shipment_document.amount, Decimal("100000.00"))

    def test_document_batch_hides_transportations_with_existing_same_document_kind(self):
        transportation = self.shipment.transportation
        ShipmentDocument.objects.create(
            transportation=transportation,
            shipment=self.shipment,
            direction=ShipmentDocument.Direction.OUTGOING,
            kind=ShipmentDocument.Kind.UPD,
            party=ShipmentDocument.Party.CUSTOMER,
            status=ShipmentDocument.Status.ISSUED,
            number="УПД-УЖЕ-ЕСТЬ",
            document_date=date.today(),
            amount=transportation.customer_amount,
            vat_amount=transportation.customer_vat_amount,
            currency=transportation.currency,
            created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("document-batch-create"),
            {
                "direction": DocumentBatch.Direction.OUTGOING,
                "owner_company": transportation.owner_company_id,
                "default_kind": ShipmentDocument.Kind.UPD,
                "currency": transportation.currency,
            },
        )
        self.assertEqual(response.status_code, 200)
        candidate_ids = {
            item["transportation"].pk
            for item in response.context["document_batch_candidates"]
        }
        self.assertNotIn(transportation.pk, candidate_ids)

    def test_transportation_executor_application_download_creates_docx(self):
        transportation = self.shipment.transportation
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("transportation-executor-application", args=[transportation.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        content = b"".join(response.streaming_content)
        document = Document(BytesIO(content))
        text_parts = [paragraph.text for paragraph in document.paragraphs]
        text_parts.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        text = "\n".join(text_parts)
        self.assertIn("ДОГОВОР-ЗАЯВКА на перевозку груза", text)

    def test_transportation_route_and_docx_show_every_stop(self):
        transportation = self.shipment.transportation
        transportation.stops.all().delete()
        for sequence, kind, city in (
            (1, TransportationStop.Kind.PICKUP, "Санкт-Петербург"),
            (2, TransportationStop.Kind.INTERMEDIATE, "Тверь"),
            (3, TransportationStop.Kind.DELIVERY, "Москва"),
        ):
            TransportationStop.objects.create(
                transportation=transportation,
                sequence=sequence,
                kind=kind,
                city=city,
                address=f"{city}, склад {sequence}",
            )

        self.assertEqual(transportation.route, "Санкт-Петербург → Тверь → Москва")
        self.client.force_login(self.user)
        response = self.client.get(transportation.get_absolute_url())
        self.assertContains(response, "Санкт-Петербург → Тверь → Москва")
        self.assertContains(response, "Промежуточная точка")
        self.assertContains(response, "Тверь")

        docx_response = self.client.get(
            reverse("transportation-executor-application", args=[transportation.pk])
        )
        content = b"".join(docx_response.streaming_content)
        document = Document(BytesIO(content))
        text = "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertIn("Санкт-Петербург → Тверь → Москва", text)
        self.assertIn("Промежуточная точка", text)
        self.assertIn("Тверь, склад 2", text)

    def test_transportation_edit_form_shows_multiple_delivery_stops(self):
        transportation = self.shipment.transportation
        transportation.stops.all().delete()
        for sequence, kind, city in (
            (1, TransportationStop.Kind.PICKUP, "Санкт-Петербург"),
            (2, TransportationStop.Kind.DELIVERY, "Москва"),
            (3, TransportationStop.Kind.DELIVERY, "Нижний Новгород"),
        ):
            TransportationStop.objects.create(
                transportation=transportation,
                sequence=sequence,
                kind=kind,
                city=city,
                address=f"{city}, склад {sequence}",
            )
        self.client.force_login(self.user)

        response = self.client.get(reverse("transportation-update", args=[transportation.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="route_stops-TOTAL_FORMS" value="3"')
        self.assertContains(response, 'name="route_stops-1-city" value="Москва"')
        self.assertContains(response, 'name="route_stops-2-city" value="Нижний Новгород"')

    def test_transportation_edit_form_renders_time_fields_as_text_masks(self):
        transportation = self.shipment.transportation
        stop = transportation.stops.order_by("sequence").first()
        planned_date = date.today()
        stop.planned_from = timezone.make_aware(
            datetime.combine(planned_date, time(9, 0))
        )
        stop.planned_to = timezone.make_aware(
            datetime.combine(planned_date, time(18, 30))
        )
        stop.save(update_fields=["planned_from", "planned_to", "updated_at"])
        self.client.force_login(self.user)

        response = self.client.get(reverse("transportation-update", args=[transportation.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'placeholder="ЧЧ:ММ"')
        self.assertContains(response, 'data-crm-time')
        self.assertContains(response, 'value="09:00"')
        self.assertContains(response, 'value="18:30"')
        self.assertNotContains(response, 'type="time"')

    def test_attached_document_download_requires_login(self):
        with tempfile.TemporaryDirectory() as media_directory:
            with override_settings(MEDIA_ROOT=media_directory):
                document_record = ShipmentDocument.objects.create(
                    shipment=self.shipment,
                    kind=ShipmentDocument.Kind.INVOICE,
                    party=ShipmentDocument.Party.CUSTOMER,
                    status=ShipmentDocument.Status.RECEIVED,
                    number="СЧ-101",
                    file=SimpleUploadedFile(
                        "invoice.pdf", b"%PDF-1.4 test invoice", "application/pdf"
                    ),
                    created_by=self.user,
                )
                download_url = reverse(
                    "shipment-document-download", args=[document_record.pk]
                )
                anonymous_response = self.client.get(download_url)
                self.assertEqual(anonymous_response.status_code, 302)
                self.client.force_login(self.user)
                response = self.client.get(download_url)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], "application/pdf")
                self.assertEqual(
                    b"".join(response.streaming_content), b"%PDF-1.4 test invoice"
                )

    def test_shipment_detail_contains_document_actions(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("shipment-detail", args=[self.shipment.pk])
        )
        self.assertContains(response, "Документы по заявке")
        self.assertContains(
            response, reverse("shipment-document-create", args=[self.shipment.pk])
        )
        self.assertContains(response, reverse("forwarding-order", args=[self.shipment.pk]))
        self.assertContains(
            response, reverse("accounting-document-create", args=[self.shipment.pk])
        )

    def test_expeditor_can_be_updated(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("expeditor-update", args=[self.company_profile.pk]),
            {
                "name": 'ООО "Новый экспедитор"',
                "short_name": 'ООО "Новый экспедитор"',
                "tax_id": "7701999999",
                "kpp": "770101001",
                "legal_address": "Москва, новый адрес",
                "default_vat_rate": CompanyProfile.VATRate.WITHOUT_VAT,
                "profit_tax_rate": "18.50",
                "is_active": "on",
            },
        )
        self.assertRedirects(response, self.company_profile.get_absolute_url())
        self.company_profile.refresh_from_db()
        self.assertEqual(self.company_profile.tax_id, "7701999999")
        self.company_profile.organization.refresh_from_db()
        self.assertEqual(
            self.company_profile.organization.profit_tax_rate, Decimal("18.50")
        )

    def test_authenticated_user_can_create_expeditor(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("expeditor-create"),
            {
                "name": 'ООО "Новая компания"',
                "short_name": 'ООО "Новая компания"',
                "tax_id": "7811999999",
                "legal_address": "Санкт-Петербург, новый адрес",
                "default_vat_rate": CompanyProfile.VATRate.WITHOUT_VAT,
                "is_active": "on",
            },
        )
        expeditor = CompanyProfile.objects.get(tax_id="7811999999")
        self.assertRedirects(response, expeditor.get_absolute_url())

    def test_dashboard_and_reports_filter_by_expeditor(self):
        second = self.create_second_expeditor()
        second_shipment = self.create_shipment_for(second)
        self.client.force_login(self.user)

        dashboard = self.client.get(
            reverse("dashboard"), {"expeditor": self.company_profile.pk}
        )
        self.assertEqual(dashboard.context["active_count"], 1)
        self.assertEqual(dashboard.context["revenue"], Decimal("100000"))
        self.assertNotContains(dashboard, second_shipment.number)

        report = self.client.get(
            reverse("reports"),
            {"currency": "RUB", "expeditor": second.pk},
        )
        self.assertEqual(report.context["shipment_count"], 1)
        self.assertEqual(report.context["totals"]["revenue"], Decimal("300000"))
        self.assertEqual(report.context["expeditor_rows"][0]["expeditor"], second)

    def test_reports_group_financials_by_expeditor(self):
        second = self.create_second_expeditor()
        self.create_shipment_for(second)
        self.client.force_login(self.user)
        response = self.client.get(reverse("reports"), {"currency": "RUB"})
        rows = {
            row["expeditor"].pk: row for row in response.context["expeditor_rows"]
        }
        self.assertEqual(rows[self.company_profile.pk]["revenue"], Decimal("100000"))
        self.assertEqual(rows[second.pk]["margin"], Decimal("50000"))

    def test_document_registry_filters_by_expeditor(self):
        second = self.create_second_expeditor()
        second_shipment = self.create_shipment_for(second)
        ShipmentDocument.objects.create(
            shipment=self.shipment,
            kind=ShipmentDocument.Kind.INVOICE,
            party=ShipmentDocument.Party.CUSTOMER,
            status=ShipmentDocument.Status.ISSUED,
            number="FIRST-COMPANY",
            created_by=self.user,
        )
        ShipmentDocument.objects.create(
            shipment=second_shipment,
            kind=ShipmentDocument.Kind.INVOICE,
            party=ShipmentDocument.Party.CUSTOMER,
            status=ShipmentDocument.Status.ISSUED,
            number="SECOND-COMPANY",
            created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("shipment-document-list"), {"expeditor": second.pk}
        )
        self.assertContains(response, "SECOND-COMPANY")
        self.assertNotContains(response, "FIRST-COMPANY")
        self.assertEqual(response.context["document_count"], 1)

    def test_accounting_document_form_prefills_shipment_and_company(self):
        other = self.create_second_expeditor()
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("accounting-document-create", args=[self.shipment.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.company_profile.tax_id)
        self.assertNotContains(response, other.tax_id)
        self.assertContains(response, self.customer.name)
        self.assertContains(response, self.shipment.number)
        self.assertContains(response, "Сформировать и скачать DOCX")
        self.assertContains(response, "Заполнение по ИНН через DaData")
        self.assertContains(response, 'data-inn-id="id_buyer_tax_id"')

    def test_all_accounting_document_forms_download_and_enter_registry(self):
        titles = {
            ShipmentDocument.Kind.INVOICE: "СЧЁТ НА ОПЛАТУ",
            ShipmentDocument.Kind.ACT: "АКТ ОКАЗАННЫХ УСЛУГ",
            ShipmentDocument.Kind.UPD: "УНИВЕРСАЛЬНЫЙ ПЕРЕДАТОЧНЫЙ ДОКУМЕНТ",
            ShipmentDocument.Kind.VAT_INVOICE: "СЧЁТ-ФАКТУРА",
        }
        self.client.force_login(self.user)
        with tempfile.TemporaryDirectory() as media_directory:
            with override_settings(MEDIA_ROOT=media_directory):
                for index, (kind, title) in enumerate(titles.items(), start=1):
                    with self.subTest(kind=kind):
                        number = f"БУХ-{index:03d}"
                        response = self.client.post(
                            reverse(
                                "accounting-document-create",
                                args=[self.shipment.pk],
                            ),
                            {
                                "kind": kind,
                                "number": number,
                                "document_date": date.today().isoformat(),
                                "contract_number": "ТЭ-01",
                                "contract_date": date.today().isoformat(),
                                "service_name": "Транспортно-экспедиционные услуги",
                                "quantity": "1.00",
                                "unit": "усл.",
                                "amount": "1220.00",
                                "currency": "RUB",
                                "vat_rate": CompanyProfile.VATRate.TWENTY_TWO,
                                "upd_status": "1",
                                "buyer_name": self.customer.name,
                                "buyer_tax_id": self.customer.tax_id,
                                "buyer_kpp": self.customer.kpp,
                                "buyer_address": self.customer.address,
                                "notes": "Тестовый документ",
                            },
                        )
                        self.assertEqual(response.status_code, 200)
                        self.assertEqual(
                            response["Content-Type"],
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                        content = b"".join(response.streaming_content)
                        document = Document(BytesIO(content))
                        text_content = "\n".join(
                            [paragraph.text for paragraph in document.paragraphs]
                            + [
                                cell.text
                                for table in document.tables
                                for row in table.rows
                                for cell in row.cells
                            ]
                        )
                        self.assertIn(title, text_content)
                        self.assertIn(self.customer.name, text_content)
                        self.assertIn("1 220,00", text_content)
                        record = ShipmentDocument.objects.get(number=number)
                        self.assertEqual(record.kind, kind)
                        self.assertEqual(record.status, ShipmentDocument.Status.ISSUED)
                        self.assertTrue(record.file.name.endswith(".docx"))
                self.assertEqual(
                    ShipmentDocument.objects.filter(shipment=self.shipment).count(),
                    4,
                )

    def test_theme_switcher_is_available_on_authenticated_pages(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-theme-choice="light"')
        self.assertContains(response, 'data-theme-choice="dark"')
        self.assertContains(response, "Светлая")
        self.assertContains(response, "Тёмная")
        self.assertContains(response, "crm-color-theme")
        self.assertContains(response, "js/theme.js")

    def test_search_is_case_insensitive_for_cyrillic(self):
        organization = Organization.objects.create(
            name="СЕЛЕКТ ЛОДЖИСТИК",
            short_name="СЕЛЕКТ ЛОДЖИСТИК",
            tax_id="7812999911",
        )
        Shipment.objects.filter(pk=self.shipment.pk).update(
            cargo_name="ПЛАСТИКОВАЯ ТАРА"
        )
        self.client.force_login(self.user)

        organizations = self.client.get(
            reverse("organization-list"),
            {"q": "селе"},
        )
        shipments = self.client.get(reverse("shipment-list"), {"q": "пласт"})

        self.assertContains(organizations, organization.name)
        self.assertContains(shipments, self.shipment.number)

    def test_payment_can_be_edited_and_deleted(self):
        payment = Payment.objects.create(
            shipment=self.shipment,
            direction=Payment.Direction.INCOME,
            amount=Decimal("10000.00"),
            payment_date=date.today(),
            reference="ПП-EDIT",
            created_by=self.user,
        )
        self.client.force_login(self.user)
        update_url = reverse(
            "payment-update",
            args=[self.shipment.pk, payment.pk],
        )
        delete_url = reverse(
            "payment-delete",
            args=[self.shipment.pk, payment.pk],
        )

        detail = self.client.get(self.shipment.get_absolute_url())
        self.assertContains(detail, update_url)
        self.assertContains(detail, delete_url)
        response = self.client.post(
            update_url,
            {
                "direction": Payment.Direction.INCOME,
                "amount": "17500.00",
                "payment_date": date.today().isoformat(),
                "method": Payment.Method.BANK,
                "reference": "ПП-EDITED",
                "notes": "Исправленная сумма",
            },
        )
        self.assertRedirects(response, self.shipment.get_absolute_url())
        payment.refresh_from_db()
        self.assertEqual(payment.amount, Decimal("17500.00"))
        self.assertEqual(payment.reference, "ПП-EDITED")
        self.assertEqual(self.shipment.received_amount, Decimal("17500.00"))

        response = self.client.post(delete_url)
        self.assertRedirects(response, self.shipment.get_absolute_url())
        self.assertFalse(Payment.objects.filter(pk=payment.pk).exists())
        self.assertEqual(self.shipment.received_amount, Decimal("0"))

    def test_delete_dependencies_link_to_related_records(self):
        Payment.objects.create(
            shipment=self.shipment,
            direction=Payment.Direction.INCOME,
            amount=Decimal("5000.00"),
            payment_date=date.today(),
            reference="ПП-LINK",
            created_by=self.user,
        )
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("shipment-delete", args=[self.shipment.pk])
        )
        payment = self.shipment.payments.get(reference="ПП-LINK")

        self.assertContains(response, payment.get_absolute_url())
        self.assertContains(response, "Поступление от клиента")
        self.assertContains(response, 'uk-icon="link"')
