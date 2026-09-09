from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from django.utils import timezone


MONTHS_GENITIVE = (
    "",
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)


def _set_run_font(run, size=10, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return run


def _date_text(value):
    if not value:
        return "«___» ____________ 20___ г."
    return f"«{value.day:02d}» {MONTHS_GENITIVE[value.month]} {value.year} г."


def _value(value, placeholder="________________________________________"):
    if value is None or value == "":
        return placeholder
    return str(value)


def _decimal_text(value):
    if value is None:
        return "________"
    return format(value, "f").rstrip("0").rstrip(".") or "0"


def _money_text(value, currency="RUB"):
    symbols = {"RUB": "₽", "RUR": "₽", "USD": "$", "EUR": "€", "CNY": "¥"}
    try:
        amount = Decimal(value or 0).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except Exception:
        amount = Decimal("0.00")
    formatted = f"{amount:,.2f}".replace(",", " ").replace(".", ",")
    return f"{formatted} {symbols.get(str(currency or '').upper(), currency or '')}".strip()


def _date_plain(value):
    return value.strftime("%d.%m.%Y") if value else "не указана"


def _datetime_window(start, end):
    if not start and not end:
        return "не указано"
    start = timezone.localtime(start) if start else None
    end = timezone.localtime(end) if end else None
    if start and end:
        if start.date() == end.date():
            return f"{start:%d.%m.%Y} г. время {start:%H:%M} {end:%H:%M}"
        return f"{start:%d.%m.%Y %H:%M} — {end:%d.%m.%Y %H:%M}"
    if start:
        return f"{start:%d.%m.%Y %H:%M}"
    return f"до {end:%d.%m.%Y %H:%M}"


def _full_address(city, address):
    if not address:
        return city
    if city and city.casefold() not in address.casefold():
        return f"{city}, {address}"
    return address


def _paragraph(document, text="", *, align=None, space_after=4, size=10):
    paragraph = document.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.0
    if text:
        _set_run_font(paragraph.add_run(text), size=size)
    return paragraph


def _section_title(document, title):
    paragraph = _paragraph(document, space_after=3)
    _set_run_font(paragraph.add_run(title), size=10.5, bold=True)
    return paragraph


def _field_row(table, label, value, *, placeholder=None):
    cells = table.add_row().cells
    cells[0].width = Cm(5.2)
    cells[1].width = Cm(12.3)
    cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    label_paragraph = cells[0].paragraphs[0]
    value_paragraph = cells[1].paragraphs[0]
    for paragraph in (label_paragraph, value_paragraph):
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
    _set_run_font(label_paragraph.add_run(label), bold=True)
    shown_value = _value(value, placeholder) if placeholder else _value(value)
    _set_run_font(value_paragraph.add_run(shown_value))
    return cells


def _details_table(document):
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    return table


def _set_cell_text(cell, text, *, bold=False, size=9, align=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if align is not None:
        paragraph.alignment = align
    paragraph.clear()
    _set_run_font(paragraph.add_run(str(text or "")), size=size, bold=bold)
    return cell


def _full_organization_name(organization):
    if not organization:
        return "не указано"
    return (
        getattr(organization, "full_name", "")
        or getattr(organization, "name_for_documents", "")
        or getattr(organization, "short_name", "")
        or str(organization)
    )


def _organization_requisites(organization):
    if not organization:
        return "реквизиты не указаны"
    parts = []
    tax_id = getattr(organization, "tax_id", "")
    kpp = getattr(organization, "kpp", "")
    ogrn = getattr(organization, "ogrn", "")
    address = (
        getattr(organization, "legal_address", "")
        or getattr(organization, "address", "")
    )
    if tax_id:
        parts.append(f"ИНН {tax_id}")
    if kpp:
        parts.append(f"КПП {kpp}")
    if ogrn:
        parts.append(f"ОГРН {ogrn}")
    if address:
        parts.append(address)
    return ", ".join(parts) if parts else "реквизиты не указаны"


def _organization_name_with_inn(organization):
    if not organization:
        return "не указан"
    tax_id = getattr(organization, "tax_id", "")
    suffix = f", ИНН {tax_id}" if tax_id else ", ИНН не указан"
    return f"{_full_organization_name(organization)}{suffix}"


def _stop_address(stop):
    if not stop:
        return "не указан"
    if stop.address:
        return _full_address(stop.city, stop.address)
    return stop.city or "не указан"


def _driver_passport_text(driver):
    if not driver:
        return "не указан"
    passport = getattr(driver, "current_passport", None)
    if passport:
        number = " ".join(
            part for part in (passport.series, passport.number) if part
        )
        details = " ".join(
            part
            for part in (
                number,
                passport.issued_by,
                f"от {_date_plain(passport.issue_date)}" if passport.issue_date else "",
            )
            if part
        )
        return details or "паспорт не указан"
    number = " ".join(
        part for part in (driver.passport_series, driver.passport_number) if part
    )
    details = " ".join(
        part
        for part in (
            number,
            driver.passport_issued_by,
            f"от {_date_plain(driver.passport_issue_date)}" if driver.passport_issue_date else "",
        )
        if part
    )
    return details or "паспорт не указан"


def _vehicle_name(vehicle):
    if not vehicle:
        return "не указано"
    title = " ".join(part for part in (vehicle.make, vehicle.model) if part).strip()
    kind = vehicle.get_kind_display() if hasattr(vehicle, "get_kind_display") else ""
    return ", ".join(part for part in (kind, title) if part) or str(vehicle)


def build_executor_transportation_application_docx(transportation):
    """Build a printable DOCX application for the selected executor."""

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.3)
    section.right_margin = Cm(1.3)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal_style.font.size = Pt(9)

    link = transportation.active_execution_link()
    assignment = transportation.active_vehicle_assignment()
    executor = link.contractor_party.organization if link else None
    contract = link.contract if link else None
    stops = list(transportation.stops.all())
    pickup_stops = [stop for stop in stops if stop.kind == stop.Kind.PICKUP]
    delivery_stops = [stop for stop in stops if stop.kind == stop.Kind.DELIVERY]
    if not pickup_stops and stops:
        pickup_stops = [stops[0]]
    if not delivery_stops and len(stops) > 1:
        delivery_stops = [stops[-1]]

    number = (
        link.instruction_number
        if link and link.instruction_number
        else transportation.number or f"рейс-{transportation.pk}"
    )
    document.core_properties.title = f"Договор-заявка {number}"
    document.core_properties.subject = "Заявка исполнителю на перевозку груза"

    title = _paragraph(document, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _set_run_font(
        title.add_run(
            f"ДОГОВОР-ЗАЯВКА на перевозку груза №{number} "
            f"от {_date_plain(transportation.document_date)} г."
        ),
        size=12,
        bold=True,
    )
    subtitle = _paragraph(document, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _set_run_font(
        subtitle.add_run(
            f"между {_full_organization_name(transportation.owner_company)} "
            f"и {_full_organization_name(executor)}"
        ),
        size=10,
    )
    contract_text = (
        f"согласно договору {contract.number} от {_date_plain(contract.contract_date)} г."
        if contract
        else "договор с исполнителем не указан"
    )
    contract_paragraph = _paragraph(document, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    _set_run_font(contract_paragraph.add_run(contract_text), size=10)

    parties_table = document.add_table(rows=2, cols=2)
    parties_table.style = "Table Grid"
    parties_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_cell_text(parties_table.cell(0, 0), "Экспедитор", bold=True)
    _set_cell_text(parties_table.cell(0, 1), "Исполнитель", bold=True)
    _set_cell_text(
        parties_table.cell(1, 0),
        f"{_full_organization_name(transportation.owner_company)}\n"
        f"{_organization_requisites(transportation.owner_company)}",
        size=8,
    )
    _set_cell_text(
        parties_table.cell(1, 1),
        f"{_full_organization_name(executor)}\n{_organization_requisites(executor)}",
        size=8,
    )

    _paragraph(document, space_after=4)
    cargo_table = document.add_table(rows=2, cols=5)
    cargo_table.style = "Table Grid"
    cargo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, heading in enumerate(("Характер груза", "Вес", "Вид упаковки", "Кол-во", "Стоимость")):
        _set_cell_text(cargo_table.cell(0, index), heading, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    package_type = transportation.package_type.name if transportation.package_type_id else ""
    quantity = []
    if transportation.package_count:
        quantity.append(f"{transportation.package_count} мест")
    if transportation.pallet_count:
        quantity.append(f"{transportation.pallet_count} паллет")
    _set_cell_text(cargo_table.cell(1, 0), transportation.cargo_name)
    _set_cell_text(
        cargo_table.cell(1, 1),
        f"{_decimal_text((transportation.weight_kg or Decimal('0')) / Decimal('1000'))} т.",
    )
    _set_cell_text(cargo_table.cell(1, 2), package_type or "не указано")
    _set_cell_text(cargo_table.cell(1, 3), ", ".join(quantity) or "не указано")
    _set_cell_text(cargo_table.cell(1, 4), "согласно ТТН")

    _paragraph(document, space_after=4)
    route_table = document.add_table(rows=0, cols=3)
    route_table.style = "Table Grid"
    route_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = route_table.add_row().cells
    _set_cell_text(header[0], "Маршрут", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    header[0].merge(header[1]).merge(header[2])
    route_row = route_table.add_row().cells
    _set_cell_text(route_row[0], transportation.route, align=WD_ALIGN_PARAGRAPH.CENTER)
    route_row[0].merge(route_row[1]).merge(route_row[2])

    def add_stop_rows(kind_title, stop_list, date_label):
        for index, stop in enumerate(stop_list, start=1):
            spacer = route_table.add_row().cells
            spacer[0].merge(spacer[1]).merge(spacer[2])
            rows = (
                (f"{kind_title} {index}", _organization_name_with_inn(stop.organization)),
                (f"Адрес {'погрузки' if stop.kind == stop.Kind.PICKUP else 'выгрузки'}", _stop_address(stop)),
                ("Контакт / комментарий", " / ".join(part for part in (stop.contact_phone, stop.contact_name, stop.instructions) if part)),
                (date_label, _datetime_window(stop.planned_from, stop.planned_to)),
            )
            for label, value in rows:
                cells = route_table.add_row().cells
                _set_cell_text(cells[0], label, bold=True)
                _set_cell_text(cells[1], value or "не указано")
                cells[1].merge(cells[2])

    add_stop_rows("Грузоотправитель", pickup_stops, "Дата погрузки")
    add_stop_rows("Грузополучатель", delivery_stops, "Дата выгрузки")

    vehicle_header = route_table.add_row().cells
    _set_cell_text(vehicle_header[0], "ТС", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(vehicle_header[1], "Водитель", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    vehicle_header[1].merge(vehicle_header[2])
    vehicle = assignment.vehicle if assignment else None
    trailer = assignment.trailer if assignment else None
    driver = assignment.driver if assignment else None
    row = route_table.add_row().cells
    _set_cell_text(row[0], _vehicle_name(vehicle))
    _set_cell_text(row[1], driver.full_name if driver else "не указан")
    row[1].merge(row[2])
    row = route_table.add_row().cells
    vehicle_numbers = " / ".join(
        part
        for part in (
            vehicle.registration_number if vehicle else "",
            trailer.registration_number if trailer else (assignment.trailer_registration_number if assignment else ""),
        )
        if part
    )
    _set_cell_text(row[0], vehicle_numbers or "не указан")
    _set_cell_text(row[1], _driver_passport_text(driver))
    row[1].merge(row[2])
    row = route_table.add_row().cells
    _set_cell_text(row[0], "-")
    _set_cell_text(row[1], f"Контактный телефон: {driver.phone}" if driver and driver.phone else "Контактный телефон: не указан")
    row[1].merge(row[2])
    row = route_table.add_row().cells
    loading = transportation.loading_method.name if transportation.loading_method_id else ""
    unloading = transportation.unloading_method.name if transportation.unloading_method_id else ""
    _set_cell_text(row[0], "Вид погрузки/выгрузки", bold=True)
    _set_cell_text(row[1], " / ".join(part for part in (loading, unloading) if part) or "не указано")
    row[1].merge(row[2])

    _paragraph(document, space_after=4)
    finance_table = document.add_table(rows=2, cols=4)
    finance_table.style = "Table Grid"
    finance_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, heading in enumerate(("Стоимость перевозки", "Предоплата", "Форма оплаты", "Срок оплаты")):
        _set_cell_text(finance_table.cell(0, index), heading, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    vat_text = (
        transportation.executor_vat_rate.name
        if transportation.executor_vat_rate_id
        else "НДС не облагается"
    )
    payment_terms = (
        contract.payment_terms
        if contract and contract.payment_terms
        else (
            f"{transportation.executor_payment_term_days} банковских дней, по оригиналам ТТН"
            if transportation.executor_payment_term_days
            else "по согласованию сторон"
        )
    )
    _set_cell_text(finance_table.cell(1, 0), _money_text(transportation.executor_amount, transportation.currency))
    _set_cell_text(finance_table.cell(1, 1), _money_text(0, transportation.currency))
    _set_cell_text(finance_table.cell(1, 2), f"Безналичный расчёт, {vat_text}")
    _set_cell_text(finance_table.cell(1, 3), payment_terms)

    clauses = (
        "Стороны несут ответственность за неисполнение либо ненадлежащее исполнение взятых на себя обязательств по настоящему договору-заявке в соответствии с условиями договора и законодательства РФ.",
        "Для исполнения договора-заявки Перевозчик вправе привлекать к перевозке третьих лиц. В этом случае Перевозчик несет ответственность за действия и бездействия третьих лиц, привлеченных к исполнению договора-заявки. Возложение исполнения обязательства на третье лицо не освобождает Перевозчика от ответственности перед Экспедитором за исполнение договора-заявки.",
        "В случае задержки подачи транспортного средства к месту погрузки, указанному в договоре-заявке Экспедитора более чем на 8 часов, Перевозчик выплачивает штраф Экспедитору в размере 1500 руб. в сутки (24 часа).",
        "Информировать Экспедитора заблаговременно о невозможности прихода транспорта в назначенное время и место погрузки, о задержке в пути следования к месту погрузки.",
        "В случае нарушения согласованных Сторонами сроков перевозки (доставки) Перевозчик выплачивает Экспедитору штраф в размере 1500 руб. за каждые сутки задержки при междугородней доставке.",
        "Перевозчик несет ответственность за несохранность груза в процессе перевозки с момента получения и подписания товарно-транспортных сопроводительных документов водителем до момента передачи груза грузополучателю, уполномоченному им лицу.",
        "Экспедитор обязуется принимать все меры для предотвращения простоя транспортного средства при погрузке/выгрузке груза.",
        "До момента заключения долгосрочного договора настоящая договор-заявка на перевозку имеет силу разового заказа.",
        "Стороны договорились, что факсовые и электронные копии настоящей договор-заявки имеют силу оригинала.",
    )
    for clause in clauses:
        _paragraph(document, clause, size=8, space_after=2)

    signatures = document.add_table(rows=3, cols=2)
    signatures.style = "Table Grid"
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_cell_text(signatures.cell(0, 0), "Экспедитор", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(signatures.cell(0, 1), "Исполнитель", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(signatures.cell(1, 0), _full_organization_name(transportation.owner_company))
    _set_cell_text(signatures.cell(1, 1), _full_organization_name(executor))
    _set_cell_text(signatures.cell(2, 0), "________________ / __________________")
    _set_cell_text(signatures.cell(2, 1), "________________ / __________________")

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def _insurance_text(value):
    if value == "yes":
        return "ДА / нет"
    if value == "no":
        return "да / НЕТ"
    return "да / нет (нужное подчеркнуть)"


def build_forwarding_order_docx(shipment, order):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.35)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal_style.font.size = Pt(10)

    document.core_properties.title = f"Экспедиторское поручение {shipment.number}"
    document.core_properties.subject = "Поручение экспедитору"

    contract_number = _value(order.contract_number, "__________")
    header = _paragraph(document, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)
    _set_run_font(header.add_run("Приложение № 3\n"), size=9)
    _set_run_font(
        header.add_run(
            f"к договору транспортной экспедиции № {contract_number}\n"
            f"от {_date_text(order.contract_date)}"
        ),
        size=9,
    )

    title = _paragraph(document, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _set_run_font(title.add_run("ПОРУЧЕНИЕ ЭКСПЕДИТОРУ"), size=14, bold=True)
    subtitle = _paragraph(document, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=9)
    _set_run_font(
        subtitle.add_run(
            f"по договору № {contract_number} от {_date_text(order.contract_date)}"
        ),
        size=10,
    )

    _section_title(document, "Грузоотправитель")
    table = _details_table(document)
    _field_row(table, "Наименование", order.shipper_name)
    _field_row(table, "ИНН", order.shipper_tax_id)
    _field_row(table, "Адрес", order.shipper_address)
    _field_row(table, "Контактное лицо", order.shipper_contact_name)
    _field_row(table, "Телефон", order.shipper_phone)

    _section_title(document, "Грузополучатель")
    table = _details_table(document)
    _field_row(table, "Наименование", order.consignee_name)
    _field_row(table, "Адрес", order.consignee_address)
    _field_row(table, "Контактное лицо", order.consignee_contact_name)
    _field_row(table, "Телефон", order.consignee_phone)

    _section_title(document, "Маршрут и время")
    table = _details_table(document)
    _field_row(
        table,
        "Адрес забора груза",
        _full_address(shipment.pickup_city, shipment.pickup_address),
    )
    _field_row(
        table,
        "Адрес доставки груза",
        _full_address(shipment.delivery_city, shipment.delivery_address),
    )
    _field_row(table, "Дата забора груза", _date_text(shipment.pickup_date))
    _field_row(table, "Плановая дата доставки", _date_text(shipment.delivery_date))
    _field_row(table, "Часы работы", order.pickup_hours)

    _section_title(document, "Сведения о грузе")
    table = _details_table(document)
    _field_row(table, "Наименование груза", shipment.cargo_name)
    _field_row(table, "Вид упаковки", order.packaging_type)
    _field_row(table, "Количество мест", order.package_count)
    _field_row(
        table,
        "Вес груза",
        f"{_decimal_text(shipment.weight_kg / 1000)} тонн",
    )
    _field_row(table, "Объём груза", f"{_decimal_text(shipment.volume_m3)} м³")
    dimensions = " × ".join(
        _decimal_text(value)
        for value in (order.cargo_length_m, order.cargo_width_m, order.cargo_height_m)
    )
    _field_row(table, "Габариты груза (Д × Ш × В)", f"{dimensions} м")
    _field_row(table, "Требуемый транспорт", shipment.vehicle_type)
    _field_row(table, "Особые условия и требования", order.special_conditions)
    _field_row(table, "Страхование груза", _insurance_text(order.cargo_insurance))
    _field_row(table, "Плательщик", order.payer)
    _field_row(table, "Место оплаты", order.payment_place)

    _section_title(document, "Важные примечания")
    _paragraph(
        document,
        "Приём поручений осуществляется с 09:00 до 16:00 с понедельника по "
        "пятницу, но не позднее чем за сутки до даты отправки груза.",
        space_after=3,
        size=9,
    )
    duties = (
        "обеспечить надлежащую упаковку груза для его сохранности;",
        "контролировать вес и количество грузовых мест;",
        "надёжно закрепить груз и правильно его разместить;",
        "убедиться, что в грузе отсутствуют запрещённые к перевозке вещества "
        "(едкие, ядовитые, самовозгорающиеся и т. п.).",
    )
    lead = _paragraph(document, space_after=1, size=9)
    _set_run_font(lead.add_run("Грузоотправитель обязан:"), size=9, bold=True)
    for duty in duties:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.65)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.0
        _set_run_font(paragraph.add_run(duty), size=9)
    _paragraph(
        document,
        "Грузоотправитель подтверждает, что ознакомлен с действующими тарифами "
        "и несёт ответственность за достоверность указанных сведений.",
        space_after=5,
        size=9,
    )

    _section_title(document, "Подписи сторон")
    signatures = document.add_table(rows=2, cols=2)
    signatures.style = "Table Grid"
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    signature_data = (
        (f"Экспедитор: {shipment.expeditor}", order.expediter_representative),
        ("Клиент (грузоотправитель)", order.client_representative),
    )
    for index, (label, representative) in enumerate(signature_data):
        cell = signatures.cell(0, index)
        paragraph = cell.paragraphs[0]
        _set_run_font(paragraph.add_run(label), bold=True)
        value_cell = signatures.cell(1, index)
        value_paragraph = value_cell.paragraphs[0]
        value_paragraph.paragraph_format.space_after = Pt(0)
        _set_run_font(
            value_paragraph.add_run(
                f"{_value(representative, '________________________')} / "
                "________________________\n(Ф. И. О., должность)                         (подпись)"
            ),
            size=9,
        )

    date_paragraph = _paragraph(document, space_after=5)
    _set_run_font(date_paragraph.add_run("Дата: "), bold=True)
    _set_run_font(date_paragraph.add_run(_date_text(order.order_date)))

    notes = _paragraph(document, space_after=1, size=8)
    _set_run_font(notes.add_run("Примечания по заполнению. "), size=8, bold=True)
    _set_run_font(
        notes.add_run(
            "При бумажном оформлении документ заполняется в двух экземплярах: оригинал передаётся "
            "экспедитору, копия остаётся у клиента. Формы экспедиторских "
            "документов утверждены приказом Минтранса России от 11.02.2008 № 23."
        ),
        size=8,
    )
    electronic = _paragraph(document, space_after=0, size=8)
    _set_run_font(
        electronic.add_run(
            "С 01.09.2026 экспедиторские документы формируются в электронном виде, "
            "кроме установленных случаев бумажного оформления. Сведения из этого "
            "проекта переносятся в применяемый оператором ЭДО формат и подписываются "
            "электронной подписью."
        ),
        size=8,
        italic=True,
    )

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def _accounting_document(title):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.35)
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal_style.font.size = Pt(10)
    document.core_properties.title = title
    return document


def _money(value):
    quantized = Decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{quantized:,.2f}".replace(",", " ").replace(".", ",")


def _vat_values(total, rate_value):
    total = Decimal(total).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if rate_value == "without_vat":
        return total, Decimal("0"), "Без НДС"
    rate = Decimal(rate_value)
    vat = (
        total * rate / (Decimal("100") + rate)
        if rate
        else Decimal("0")
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return total - vat, vat, f"{rate_value}%"


def _party_details(name, tax_id, kpp, address):
    tax_details = f"ИНН {tax_id or '__________'}"
    if kpp:
        tax_details += f", КПП {kpp}"
    return f"{name}, {tax_details}, адрес: {address or '____________________'}"


def _add_title(document, title, number, document_date):
    paragraph = _paragraph(document, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=9)
    _set_run_font(
        paragraph.add_run(f"{title} № {number} от {_date_text(document_date)}"),
        size=14,
        bold=True,
    )


def _set_table_cell(cell, text, *, bold=False, size=9, align=None):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    if align is not None:
        paragraph.alignment = align
    _set_run_font(paragraph.add_run(str(text)), size=size, bold=bold)


def _add_simple_item_table(document, data, net_amount, vat_amount, vat_label):
    table = document.add_table(rows=2, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headings = ("№", "Наименование услуги", "Кол-во", "Ед.", "Цена", "Сумма")
    for index, heading in enumerate(headings):
        _set_table_cell(
            table.cell(0, index), heading, bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER
        )
    quantity = Decimal(data["quantity"])
    unit_price = (net_amount / quantity).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )
    values = (
        "1",
        data["service_name"],
        _decimal_text(quantity),
        data["unit"],
        _money(unit_price),
        _money(net_amount),
    )
    for index, value in enumerate(values):
        _set_table_cell(table.cell(1, index), value, size=8)
    total_paragraph = _paragraph(document, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=1)
    _set_run_font(total_paragraph.add_run(f"Итого без налога: {_money(net_amount)}"), bold=True)
    vat_paragraph = _paragraph(document, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=1)
    vat_text = "Без НДС" if vat_label == "Без НДС" else f"НДС {vat_label}: {_money(vat_amount)}"
    _set_run_font(vat_paragraph.add_run(vat_text), bold=True)
    final_paragraph = _paragraph(document, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=7)
    _set_run_font(
        final_paragraph.add_run(f"Всего к оплате: {_money(data['amount'])} {data['currency']}"),
        size=11,
        bold=True,
    )


def _add_tax_item_table(document, data, net_amount, vat_amount, vat_label):
    table = document.add_table(rows=2, cols=9)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headings = (
        "№",
        "Наименование работ, услуг",
        "Ед.",
        "Кол-во",
        "Цена без НДС",
        "Стоимость без НДС",
        "Ставка",
        "Сумма НДС",
        "Всего",
    )
    for index, heading in enumerate(headings):
        _set_table_cell(
            table.cell(0, index), heading, bold=True, size=6.5,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    quantity = Decimal(data["quantity"])
    unit_price = (net_amount / quantity).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )
    values = (
        "1",
        data["service_name"],
        data["unit"],
        _decimal_text(quantity),
        _money(unit_price),
        _money(net_amount),
        vat_label,
        "—" if vat_label == "Без НДС" else _money(vat_amount),
        _money(data["amount"]),
    )
    for index, value in enumerate(values):
        _set_table_cell(table.cell(1, index), value, size=6.8)


def _add_signatures(document, profile, *, include_buyer=True):
    table = document.add_table(rows=2, cols=2 if include_buyer else 1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Исполнитель"]
    values = [profile.director_name]
    if include_buyer:
        headers.append("Заказчик")
        values.append("")
    for index, header in enumerate(headers):
        _set_table_cell(table.cell(0, index), header, bold=True)
        representative = values[index] or "________________________"
        _set_table_cell(
            table.cell(1, index),
            f"{representative} / ________________________\n"
            "               (Ф. И. О.)                         (подпись)",
            size=8,
        )


def build_accounting_document_docx(shipment, profile, data):
    kind = data["kind"]
    titles = {
        "invoice": "СЧЁТ НА ОПЛАТУ",
        "act": "АКТ ОКАЗАННЫХ УСЛУГ",
        "upd": "УНИВЕРСАЛЬНЫЙ ПЕРЕДАТОЧНЫЙ ДОКУМЕНТ",
        "vat_invoice": "СЧЁТ-ФАКТУРА",
    }
    title = titles[kind]
    document = _accounting_document(
        f"{title.title()} {data['number']} · {shipment.number}"
    )
    net_amount, vat_amount, vat_label = _vat_values(
        data["amount"], data["vat_rate"]
    )
    currency_code = {"RUB": "643", "USD": "840", "EUR": "978"}[
        data["currency"]
    ]
    seller = _party_details(profile.name, profile.tax_id, profile.kpp, profile.legal_address)
    buyer = _party_details(
        data["buyer_name"], data["buyer_tax_id"], data["buyer_kpp"], data["buyer_address"]
    )

    if kind == "invoice":
        bank_table = _details_table(document)
        _field_row(bank_table, "Банк получателя", profile.bank_name)
        _field_row(bank_table, "БИК", profile.bik)
        _field_row(bank_table, "Корр. счёт", profile.correspondent_account)
        _field_row(bank_table, "Получатель", seller)
        _field_row(bank_table, "Расчётный счёт", profile.settlement_account)
        _add_title(document, title, data["number"], data["document_date"])
        parties = _details_table(document)
        _field_row(parties, "Поставщик", seller)
        _field_row(parties, "Покупатель", buyer)
        _add_simple_item_table(document, data, net_amount, vat_amount, vat_label)
        _paragraph(
            document,
            f"Основание: договор № {_value(data['contract_number'], '________')} "
            f"от {_date_text(data['contract_date'])}; заявка {shipment.number}.",
            space_after=8,
        )
        _add_signatures(document, profile, include_buyer=False)

    elif kind == "act":
        _add_title(document, title, data["number"], data["document_date"])
        details = _details_table(document)
        _field_row(details, "Исполнитель", seller)
        _field_row(details, "Заказчик", buyer)
        _field_row(
            details,
            "Основание",
            f"Договор № {_value(data['contract_number'], '________')} "
            f"от {_date_text(data['contract_date'])}, заявка {shipment.number}",
        )
        _add_simple_item_table(document, data, net_amount, vat_amount, vat_label)
        _paragraph(
            document,
            "Услуги оказаны в полном объёме и в установленный срок. Заказчик "
            "претензий по объёму, качеству и срокам оказания услуг не имеет.",
            space_after=9,
        )
        _add_signatures(document, profile)

    else:
        if kind == "upd":
            function_text = (
                "счёт-фактура и передаточный документ"
                if data["upd_status"] == "1"
                else "передаточный документ"
            )
            function_paragraph = _paragraph(
                document, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2
            )
            _set_run_font(
                function_paragraph.add_run(
                    f"Статус: {data['upd_status']} — {function_text}"
                ),
                size=9,
                bold=True,
            )
        _add_title(document, title, data["number"], data["document_date"])
        details = _details_table(document)
        _field_row(details, "Продавец", seller)
        _field_row(details, "Покупатель", buyer)
        _field_row(
            details,
            "Валюта",
            f"{data['currency']} (код {currency_code})",
        )
        _field_row(
            details,
            "Основание передачи",
            f"Договор № {_value(data['contract_number'], '________')} "
            f"от {_date_text(data['contract_date'])}, заявка {shipment.number}",
        )
        _add_tax_item_table(document, data, net_amount, vat_amount, vat_label)
        totals = _paragraph(document, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=8)
        vat_text = "Без НДС" if vat_label == "Без НДС" else f"в том числе НДС {_money(vat_amount)}"
        _set_run_font(
            totals.add_run(
                f"Всего: {_money(data['amount'])} {data['currency']}, {vat_text}."
            ),
            bold=True,
        )
        if kind == "upd":
            _paragraph(
                document,
                "Содержание факта хозяйственной жизни: услуги оказаны по заявке "
                f"{shipment.number}. Дата оказания услуг: {_date_text(shipment.delivery_date)}.",
                space_after=8,
            )
        _add_signatures(document, profile, include_buyer=kind == "upd")

    if data.get("notes"):
        note = _paragraph(document, space_after=4)
        _set_run_font(note.add_run("Дополнительные сведения: "), bold=True)
        _set_run_font(note.add_run(data["notes"]))
    warning = _paragraph(document, space_after=0, size=8)
    _set_run_font(
        warning.add_run(
            "Редактируемый проект документа. Перед подписанием необходимо проверить "
            "реквизиты, налоговую ставку и применимость формы. Для юридически значимого "
            "электронного обмена используется установленный формат оператора ЭДО."
        ),
        size=8,
        italic=True,
    )

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream
